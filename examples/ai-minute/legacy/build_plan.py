# -*- coding: utf-8 -*-
"""Build the AI Minute master plan workbook (Vietnamese).

Schedule model: one person per role (lead, developer, QA). Each role works its
level-3 packages one after another ("Nối tiếp"); a package also waits for the
packages it depends on, whatever role owns them. A few packages run alongside
others ("Kéo dài", "Bảo hành"); finished work keeps its real dates ("Cố định").
All dates, gates, Gantt bars and load figures are Excel formulas.
"""

import math
import sys
from datetime import date, timedelta

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import plan_data as d

OUT = sys.argv[1]
HOL = [h[0] for h in d.HOLIDAYS]
ROLE_KEYS = list(d.ROLES)
ROLE_NAME = {k: v[0] for k, v in d.ROLES.items()}

# ---------------------------------------------------------------- rows and chains
rows = []
for ph in d.PHASES:
    rows.append(("L1", ph, ph[0]))
    for item in d.W[ph[1]]:
        rows.append((item[0], item, ph[0]))
L3ROWS = [it for kind, it, gd in rows if kind == "L3"]
BYCODE = {it[1]: it for it in L3ROWS}


def mode_of(code):
    return d.SCHED.get(code, ("seq",))


chains = {k: [] for k in ROLE_KEYS}
for it in L3ROWS:
    if mode_of(it[1])[0] == "seq":
        chains[it[9]].append(it[1])
for code, after in d.CHAIN_AFTER.items():
    ch = chains[BYCODE[code][9]]
    ch.remove(code)
    ch.insert(ch.index(after) + 1, code)
PREV = {}
for role, ch in chains.items():
    for i, code in enumerate(ch):
        PREV[code] = ch[i - 1] if i else None


# ---------------------------------------------------------------- python mirror of the schedule
def workday(day, n):
    cur, step, k = day, (1 if n > 0 else -1), abs(n)
    while k:
        cur += timedelta(days=step)
        if cur.weekday() < 5 and cur not in HOL:
            k -= 1
    return cur


def edate_minus1(day, months):
    y, m = day.year + (day.month - 1 + months) // 12, (day.month - 1 + months) % 12 + 1
    last = [31, 29 if y % 4 == 0 else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return date(y, m, min(day.day, last)) - timedelta(days=1)


def simulate():
    S, E = {}, {}
    for it in L3ROWS:
        if mode_of(it[1])[0] == "fixed":
            S[it[1]], E[it[1]] = it[6], it[7]
    seq = [it for it in L3ROWS if mode_of(it[1])[0] == "seq"]
    for _ in range(300):
        changed = False
        for it in seq:
            code = it[1]
            base = [d.PROJECT_START - timedelta(days=1)]
            if PREV[code]:
                base.append(E.get(PREV[code], date(1900, 1, 1)))
            base += [E.get(dep, date(1900, 1, 1)) for dep in it[10]]
            dur = max(1, math.floor(it[8] * (1 - it[13]) / d.ROLES[it[9]][1] + 0.5 + 1e-9))
            s_ = workday(max(base), 1)
            e_ = workday(s_, dur - 1)
            if S.get(code) != s_ or E.get(code) != e_:
                S[code], E[code], changed = s_, e_, True
        if not changed:
            break
    for it in L3ROWS:
        m = mode_of(it[1])
        if m[0] == "span":
            S[it[1]], E[it[1]] = S[m[1]], E[m[2]]
        elif m[0] == "warranty":
            S[it[1]] = workday(E["8.3.3"], 1)
            E[it[1]] = edate_minus1(S[it[1]], 12)
    for it in L3ROWS:
        if mode_of(it[1])[0] == "warranty_end":
            E[it[1]] = E["9.1.1"]
            S[it[1]] = workday(E[it[1]], -4)
    return S, E


SIM_S, SIM_E = simulate()


def group_end(code):
    return max(SIM_E[c] for c in SIM_E if c == code or c.startswith(code + "."))


END_BUILD = SIM_E["8.3.3"]
END_ALL = SIM_E["9.1.1"]

# ---------------------------------------------------------------- styles
FONT = "Arial"
NAVY, NAVY2, NAVY3 = "1E3A8A", "5B7BC4", "B7C7EA"
L1FILL, L2FILL = "D6E0F5", "EEF2F7"
thin = Side(style="thin", color="C9D3E0")
hair = Side(style="hair", color="E2E8F0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE, GREEN, INK = "0000FF", "008000", "0F1B2D"
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
VC = Alignment(wrap_text=True, vertical="center")
DATE_FMT = "dd/mm/yyyy"


def F(bold=False, color=INK, size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def fill(h):
    return PatternFill("solid", start_color=h, end_color=h)


def title(ws, text, sub, last_col, legend=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, text).font = F(True, NAVY, 14)
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, sub).font = F(False, "4A5A70", 9, True)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 28
    if legend:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
        ws.cell(3, 1, legend).font = F(False, "4A5A70", 9)


def header(ws, row, labels, widths=None, height=32, col0=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row, col0 + i, lab)
        c.font = F(True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = CENTER
        c.border = BOX
    ws.row_dimensions[row].height = height
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[CL(col0 + i)].width = w


def cell(ws, r, c, v, fmt=None, bold=False, color=INK, align=None, bg=None, border=True):
    x = ws.cell(r, c, v)
    x.font = F(bold, color)
    if fmt:
        x.number_format = fmt
    x.alignment = align or WRAP
    if bg:
        x.fill = fill(bg)
    if border:
        x.border = BOX
    return x


def page_setup(ws, title_rows, portrait=False):
    ws.page_setup.orientation = "portrait" if portrait else "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.center.text = "Trang &P / &N"
    ws.oddFooter.left.text = "AI Minute – Kế hoạch tổng thể v2.0"
    ws.page_margins.left = ws.page_margins.right = 0.4


def status_cf(ws, ref, top_left):
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'{top_left}="Hoàn thành"'], fill=fill("E7F5EC"), font=Font(name=FONT, color="15803D", bold=True)))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'{top_left}="Đang thực hiện"'], fill=fill("E3F1F9"), font=Font(name=FONT, color="0369A1", bold=True)))


wb = Workbook()
wb.calculation.fullCalcOnLoad = True
S_TQ = wb.active
S_TQ.title = "Tổng quan"
S_GD = wb.create_sheet("Giai đoạn & Phạm vi")
S_W = wb.create_sheet("WBS")
S_G = wb.create_sheet("Gantt tuần")
S_M = wb.create_sheet("Lộ trình tháng")
S_C = wb.create_sheet("Mốc & Cổng")
S_P = wb.create_sheet("Phạm vi tính năng")
S_N = wb.create_sheet("Nguồn lực")
S_R = wb.create_sheet("Rủi ro")
S_A = wb.create_sheet("Giả định & Ràng buộc")
S_H = wb.create_sheet("Ngày nghỉ")
S_D = wb.create_sheet("Danh mục")

# ---------------------------------------------------------------- Danh mục
title(S_D, "DANH MỤC GIÁ TRỊ", "Dùng cho danh sách chọn ở các sheet khác. Có thể bổ sung giá trị vào cuối mỗi cột.", 8)
header(S_D, 4, ["Trạng thái công việc", "Ưu tiên", "Ý nghĩa ưu tiên", "Phát hành", "Trạng thái cổng", "Trạng thái rủi ro", "Cách lập lịch", "Ý nghĩa cách lập lịch"],
       [22, 10, 30, 14, 20, 20, 16, 64])
lists = {
    1: ["Chưa bắt đầu", "Đang thực hiện", "Hoàn thành", "Tạm dừng", "Hủy"],
    2: ["P0", "P1", "P2"],
    3: ["Bắt buộc để bàn giao", "Nên có", "Có thể bỏ, cắt đầu tiên"],
    4: d.RELEASES,
    5: ["Chưa đến hạn", "Đạt", "Đạt có điều kiện", "Không đạt"],
    6: ["Đang theo dõi", "Đã xảy ra", "Đã đóng"],
    7: ["Nối tiếp", "Kéo dài", "Bảo hành", "Cố định"],
    8: ["Bắt đầu sau gói trước của cùng vai trò và sau mọi gói phụ thuộc; thời lượng = công còn lại ÷ tỷ lệ thời gian chuyên môn của vai trò, làm tròn đến ngày gần nhất (tối thiểu 1 ngày)",
        "Chạy song song trong khoảng thời gian của các gói khác (theo dõi, quản lý, hỗ trợ)",
        "Trải đều trong 12 tháng bảo hành sau nghiệm thu tổng thể",
        "Ngày thực tế đã diễn ra, không lập lịch lại"],
}
for col, vals in lists.items():
    for i, v in enumerate(vals):
        cell(S_D, 5 + i, col, v)
for nm, col, n in [("DM_TrangThai", "A", 5), ("DM_UuTien", "B", 3), ("DM_PhatHanh", "D", len(d.RELEASES)), ("DM_Cong", "E", 4), ("DM_RuiRo", "F", 3)]:
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'Danh mục'!${col}$5:${col}${4 + n}")
S_D.freeze_panes = "A5"

# ---------------------------------------------------------------- Ngày nghỉ
title(S_H, "NGÀY NGHỈ LỄ (loại khỏi ngày làm việc khi lập lịch)",
      "Ngày từ năm 2027 là dự kiến; cập nhật theo thông báo chính thức. Có thể thêm dòng mới trong vùng A5:A80.", 4)
header(S_H, 4, ["Ngày", "Thứ", "Tên ngày nghỉ", "Tình trạng"], [14, 8, 50, 14])
for i, (dt, name, st) in enumerate(d.HOLIDAYS):
    r = 5 + i
    cell(S_H, r, 1, dt, DATE_FMT, color=BLUE)
    cell(S_H, r, 2, f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r}),"CN","T2","T3","T4","T5","T6","T7"))', align=CENTER)
    cell(S_H, r, 3, name, color=BLUE)
    cell(S_H, r, 4, st, color=BLUE)
wb.defined_names["NgayNghi"] = DefinedName("NgayNghi", attr_text="'Ngày nghỉ'!$A$5:$A$80")
S_H.freeze_panes = "A5"

# ---------------------------------------------------------------- Nguồn lực: roles and parameters
title(S_N, "NGUỒN LỰC, NĂNG LỰC VÀ LẬP LỊCH",
      "Mỗi vai trò 1 người. Ô nền vàng là tham số: thay đổi sẽ lập lại toàn bộ tiến độ trong WBS, Gantt và các mốc.", 10)
header(S_N, 4, ["Mã", "Vai trò", "Số người", "Tỷ lệ thời gian chuyên môn", "Trách nhiệm"], None, 30)
for i, k in enumerate(ROLE_KEYS):
    r = 5 + i
    cell(S_N, r, 1, k, bold=True, align=CENTER)
    cell(S_N, r, 2, ROLE_NAME[k], bold=True, align=VC)
    cell(S_N, r, 3, 1, "0", align=CENTER)
    cell(S_N, r, 4, d.ROLES[k][1], "0%", color=BLUE, align=CENTER, bg="FFFF00")
    cell(S_N, r, 5, d.ROLES[k][2], align=VC)
    S_N.row_dimensions[r].height = 30
RL0, RL1 = 5, 4 + len(ROLE_KEYS)
wb.defined_names["DM_VaiTro"] = DefinedName("DM_VaiTro", attr_text=f"'Nguồn lực'!$B${RL0}:$B${RL1}")
PR = RL1 + 2
cell(S_N, PR, 2, "Ngày bắt đầu lập lịch", bold=True, bg=L2FILL)
cell(S_N, PR, 4, d.PROJECT_START, DATE_FMT, color=BLUE, align=CENTER, bg="FFFF00")
cell(S_N, PR + 1, 2, "Mốc nghiệm thu MVP mục tiêu", bold=True, bg=L2FILL)
cell(S_N, PR + 1, 4, d.MVP_TARGET, DATE_FMT, color=BLUE, align=CENTER, bg="FFFF00")
START = f"'Nguồn lực'!$D${PR}"
TARGET = f"'Nguồn lực'!$D${PR + 1}"


def FOCUS(r):
    return f"INDEX('Nguồn lực'!$D${RL0}:$D${RL1},MATCH($N{r},'Nguồn lực'!$B${RL0}:$B${RL1},0))"


# ---------------------------------------------------------------- WBS
WCOLS = ["Mã WBS", "Cấp", "Hạng mục công việc", "Sản phẩm bàn giao / Mô tả", "Giai đoạn", "Tham chiếu", "Ưu tiên",
         "Bắt đầu", "Kết thúc", "Thời lượng (ngày LV)", "Công (ngày)", "Cách lập lịch", "Phụ thuộc (mã WBS)", "Phụ trách",
         "Phối hợp", "Trạng thái", "% hoàn thành", "Ghi chú", "Nhóm"]
WW = [9, 6, 46, 38, 8, 16, 8, 11, 11, 10, 9, 11, 16, 15, 14, 15, 10, 26, 7]
title(S_W, "CẤU TRÚC PHÂN RÃ CÔNG VIỆC (WBS) VÀ TIẾN ĐỘ – DỰ ÁN AI MINUTE",
      "Phiên bản 2.0 · 25/09/2026 · MVP web trước, ứng dụng desktop và mobile ở GĐ5. Công = ngày công ước lượng, đã gồm thiết kế chi tiết, kiểm thử và tài liệu. "
      "Tiến độ lập tự động: mỗi vai trò làm lần lượt các gói của mình; một gói bắt đầu khi các gói phụ thuộc đã xong.",
      len(WCOLS),
      "Chữ xanh dương: số liệu nhập (chỉ ở dòng cấp 3) · Chữ đen: công thức · Dòng tô đậm: tổng hợp giai đoạn (cấp 1) và nhóm (cấp 2) · "
      "Sửa Công, % hoàn thành hoặc tỷ lệ thời gian (sheet Nguồn lực) để lập lại lịch · Thứ tự và phụ thuộc được thiết lập khi lập kế hoạch · Bấm +/− bên trái để thu gọn.")
header(S_W, 4, WCOLS, WW, 36)
FIRST = 5
LAST = FIRST + len(rows) - 1
wbs_row = {it[1]: FIRST + i for i, (kind, it, gd) in enumerate(rows)}
block_end = {}
for i, (kind, it, gd) in enumerate(rows):
    if kind == "L3":
        continue
    j = i + 1
    while j < len(rows) and not (rows[j][0] == "L1" or (kind == "L2" and rows[j][0] == "L2")):
        j += 1
    block_end[i] = FIRST + j - 1
MODE_LABEL = {"fixed": "Cố định", "seq": "Nối tiếp", "span": "Kéo dài", "warranty": "Bảo hành", "warranty_end": "Bảo hành"}
group = None
for i, (kind, it, gd) in enumerate(rows):
    r = FIRST + i
    lvl = {"L1": 1, "L2": 2, "L3": 3}[kind]
    bg = {1: L1FILL, 2: L2FILL, 3: None}[lvl]
    code, name, desc = it[1], it[2], it[3]
    if kind == "L1":
        group = None
    elif kind == "L2":
        group = code
    bold = lvl < 3
    cell(S_W, r, 1, code, bold=bold, bg=bg)
    cell(S_W, r, 2, lvl, bold=bold, bg=bg, align=CENTER)
    cell(S_W, r, 3, name, bold=bold, bg=bg).alignment = Alignment(wrap_text=True, vertical="top", indent=lvl - 1)
    cell(S_W, r, 4, desc, bg=bg)
    cell(S_W, r, 5, gd, bold=bold, bg=bg, align=CENTER)
    cell(S_W, r, 19, "" if lvl == 1 else (code if lvl == 2 else group), bg=bg)
    if lvl == 3:
        ref, prio, st, en, eff, role, deps, coop, status, pct, note = it[4:15]
        mode = mode_of(code)
        cell(S_W, r, 6, ref, color=BLUE)
        cell(S_W, r, 7, prio, color=BLUE, align=CENTER)
        cell(S_W, r, 14, ROLE_NAME[role], color=BLUE)
        if mode[0] == "fixed":
            cell(S_W, r, 8, st, DATE_FMT, color=BLUE, align=CENTER)
            cell(S_W, r, 9, en, DATE_FMT, color=BLUE, align=CENTER)
            cell(S_W, r, 10, f"=NETWORKDAYS(H{r},I{r},NgayNghi)", "0", align=CENTER)
        elif mode[0] == "seq":
            refs = [f"{START}-1"]
            if PREV[code]:
                refs.append(f"I{wbs_row[PREV[code]]}")
            refs += [f"I{wbs_row[x]}" for x in deps]
            cell(S_W, r, 8, f"=WORKDAY(MAX({','.join(refs)}),1,NgayNghi)", DATE_FMT, align=CENTER)
            cell(S_W, r, 10, f"=MAX(1,ROUND(K{r}*(1-Q{r})/{FOCUS(r)},0))", "0", align=CENTER)
            cell(S_W, r, 9, f"=WORKDAY(H{r},J{r}-1,NgayNghi)", DATE_FMT, align=CENTER)
        elif mode[0] == "span":
            cell(S_W, r, 8, f"=H{wbs_row[mode[1]]}", DATE_FMT, align=CENTER)
            cell(S_W, r, 9, f"=I{wbs_row[mode[2]]}", DATE_FMT, align=CENTER)
            cell(S_W, r, 10, f"=NETWORKDAYS(H{r},I{r},NgayNghi)", "0", align=CENTER)
        elif mode[0] == "warranty":
            cell(S_W, r, 8, f"=WORKDAY(I{wbs_row['8.3.3']},1,NgayNghi)", DATE_FMT, align=CENTER)
            cell(S_W, r, 9, f"=EDATE(H{r},12)-1", DATE_FMT, align=CENTER)
            cell(S_W, r, 10, f"=NETWORKDAYS(H{r},I{r},NgayNghi)", "0", align=CENTER)
        else:
            cell(S_W, r, 9, f"=I{wbs_row['9.1.1']}", DATE_FMT, align=CENTER)
            cell(S_W, r, 8, f"=WORKDAY(I{r},-4,NgayNghi)", DATE_FMT, align=CENTER)
            cell(S_W, r, 10, f"=NETWORKDAYS(H{r},I{r},NgayNghi)", "0", align=CENTER)
        if code in d.QL_RATE:
            cell(S_W, r, 11, f"=ROUND(J{r}/5*{d.QL_RATE[code]},1)", "0.0#", align=CENTER)
            note = f"{str(d.QL_RATE[code]).replace('.', ',')} ngày/tuần trong khoảng thời gian"
        else:
            cell(S_W, r, 11, eff, "0.0#", color=BLUE, align=CENTER)
        cell(S_W, r, 12, MODE_LABEL[mode[0]], align=CENTER)
        cell(S_W, r, 13, "; ".join(deps) if deps else "—")
        cell(S_W, r, 15, coop, color=BLUE)
        cell(S_W, r, 16, status, color=BLUE)
        cell(S_W, r, 17, pct, "0%", color=BLUE, align=CENTER)
        cell(S_W, r, 18, note, color=BLUE)
    else:
        a, b = r + 1, block_end[i]
        blk = lambda col: f"${col}${a}:${col}${b}"
        for c_ in (6, 7, 12, 13, 15, 18):
            cell(S_W, r, c_, "", bg=bg)
        cell(S_W, r, 8, f"=MIN({blk('H')})", DATE_FMT, bold=True, bg=bg, align=CENTER)
        cell(S_W, r, 9, f"=MAX({blk('I')})", DATE_FMT, bold=True, bg=bg, align=CENTER)
        cell(S_W, r, 10, f"=NETWORKDAYS(H{r},I{r},NgayNghi)", "0", bold=True, bg=bg, align=CENTER)
        cell(S_W, r, 11, f"=SUMIFS({blk('K')},{blk('B')},3)", "0.0#", bold=True, bg=bg, align=CENTER)
        cell(S_W, r, 14, "Cả đội", bg=bg)
        cell(S_W, r, 16, f'=IF(Q{r}>=1,"Hoàn thành",IF(Q{r}>0,"Đang thực hiện","Chưa bắt đầu"))', bold=True, bg=bg)
        cell(S_W, r, 17, f"=IFERROR(SUMPRODUCT(({blk('B')}=3)*{blk('K')}*{blk('Q')})/K{r},0)", "0%", bold=True, bg=bg, align=CENTER)
    S_W.row_dimensions[r].outlineLevel = lvl - 1
S_W.sheet_properties.outlinePr.summaryBelow = False
S_W.column_dimensions["S"].hidden = True
S_W.freeze_panes = "D5"
S_W.auto_filter.ref = f"A4:S{LAST}"
rng = lambda col: f"${col}${FIRST}:${col}${LAST}"
W_TOT = LAST + 2
cell(S_W, W_TOT, 3, "TỔNG TOÀN DỰ ÁN (cộng các gói công việc cấp 3)", bold=True, bg=L1FILL)
for c_ in (1, 2, 4, 5, 6, 7, 10, 12, 13, 14, 15, 16, 18):
    cell(S_W, W_TOT, c_, "", bg=L1FILL)
cell(S_W, W_TOT, 8, f"=MIN({rng('H')})", DATE_FMT, True, bg=L1FILL, align=CENTER)
cell(S_W, W_TOT, 9, f"=MAX({rng('I')})", DATE_FMT, True, bg=L1FILL, align=CENTER)
cell(S_W, W_TOT, 11, f"=SUMIFS({rng('K')},{rng('B')},3)", "0.0#", True, bg=L1FILL, align=CENTER)
cell(S_W, W_TOT, 17, f"=IFERROR(SUMPRODUCT(({rng('B')}=3)*{rng('K')}*{rng('Q')})/K{W_TOT},0)", "0%", True, bg=L1FILL, align=CENTER)
dv = DataValidation(type="list", formula1="=DM_TrangThai", allow_blank=True)
dv2 = DataValidation(type="list", formula1="=DM_UuTien", allow_blank=True)
dv3 = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True, error="Nhập tỉ lệ từ 0% đến 100%", errorTitle="% hoàn thành")
dv4 = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, error="Công phải là số không âm", errorTitle="Công")
dv5 = DataValidation(type="list", formula1="=DM_VaiTro", allow_blank=False)
for i, (kind, it, gd) in enumerate(rows):
    if kind == "L3":
        r = FIRST + i
        dv.add(f"P{r}")
        dv2.add(f"G{r}")
        dv3.add(f"Q{r}")
        dv5.add(f"N{r}")
        if it[1] not in d.QL_RATE:
            dv4.add(f"K{r}")
for v in (dv, dv2, dv3, dv4, dv5):
    S_W.add_data_validation(v)
status_cf(S_W, f"P{FIRST}:P{LAST}", f"P{FIRST}")
S_W.conditional_formatting.add(f"I{FIRST}:I{LAST}", FormulaRule(formula=[f'AND($B{FIRST}=3,I{FIRST}<TODAY(),$P{FIRST}<>"Hoàn thành",$P{FIRST}<>"Hủy")'], fill=fill("FBE9E9"), font=Font(name=FONT, color="B91C1C", bold=True)))
page_setup(S_W, "4:4")

# ---------------------------------------------------------------- Mốc & Cổng
CCOLS = ["Mã", "Mốc / Cổng", "Giai đoạn", "Hạng mục WBS", "Ngày kế hoạch", "Tiêu chí đạt", "Bằng chứng / Sản phẩm", "Phê duyệt", "Trạng thái", "Ngày thực tế", "Chênh lệch (ngày LV)"]
title(S_C, "MỐC VÀ CỔNG KIỂM SOÁT",
      "Ngày kế hoạch = ngày kết thúc của hạng mục WBS tương ứng (tự cập nhật khi tiến độ thay đổi). Nhập Trạng thái và Ngày thực tế khi họp cổng; chênh lệch dương là trễ.", len(CCOLS))
header(S_C, 4, CCOLS, [7, 32, 9, 10, 13, 50, 32, 22, 16, 13, 12])
gate_row = {}
for i, g in enumerate(d.GATES):
    r = 5 + i
    gate_row[g[0]] = r
    cell(S_C, r, 1, g[0], bold=True, align=CENTER)
    cell(S_C, r, 2, g[1], bold=True)
    cell(S_C, r, 3, g[2], align=CENTER)
    cell(S_C, r, 4, d.GATE_WBS[g[0]], color=BLUE, align=CENTER)
    cell(S_C, r, 5, f"=INDEX(WBS!$I:$I,MATCH($D{r},WBS!$A:$A,0))", DATE_FMT, bold=True, color=GREEN, align=CENTER)
    cell(S_C, r, 6, g[3])
    cell(S_C, r, 7, g[4])
    cell(S_C, r, 8, g[5])
    cell(S_C, r, 9, "Chưa đến hạn", color=BLUE)
    cell(S_C, r, 10, None, DATE_FMT, color=BLUE, align=CENTER)
    cell(S_C, r, 11, f'=IF(J{r}="","",IF(J{r}>=E{r},NETWORKDAYS(E{r},J{r},NgayNghi)-1,-(NETWORKDAYS(J{r},E{r},NgayNghi)-1)))', "0", align=CENTER)
    S_C.row_dimensions[r].height = 45
C_LAST = 4 + len(d.GATES)
dvc = DataValidation(type="list", formula1="=DM_Cong", allow_blank=True)
dvc.add(f"I5:I{C_LAST}")
S_C.add_data_validation(dvc)
for val, bg_, fg_ in [("Đạt", "E7F5EC", "15803D"), ("Đạt có điều kiện", "FDF1E3", "B45309"), ("Không đạt", "FBE9E9", "B91C1C")]:
    S_C.conditional_formatting.add(f"I5:I{C_LAST}", FormulaRule(formula=[f'I5="{val}"'], fill=fill(bg_), font=Font(name=FONT, color=fg_, bold=True)))
S_C.freeze_panes = "C5"
page_setup(S_C, "4:4")


def GATE_DATE(g):
    return f"'Mốc & Cổng'!$E${gate_row[g]}"


# ---------------------------------------------------------------- Giai đoạn & Phạm vi
GCOLS = ["Mã GĐ", "Giai đoạn", "Bắt đầu", "Kết thúc", "Số tuần", "Công (ngày)", "Cổng", "Ngày cổng", "Mục tiêu", "Phạm vi thực hiện",
         "Ngoài phạm vi", "Sản phẩm bàn giao", "Tiêu chí kết thúc (cổng)", "Rủi ro chính", "Trạng thái", "% hoàn thành"]
title(S_GD, "GIAI ĐOẠN, PHẠM VI VÀ TIÊU CHÍ KẾT THÚC",
      "Toàn bộ vòng đời dự án. MVP là bản web; ứng dụng Windows và Android (GĐ5) là cách dùng được khuyến nghị của sản phẩm cuối. Ngày, công, cổng và trạng thái lấy tự động từ WBS.", len(GCOLS))
header(S_GD, 4, GCOLS, [8, 22, 11, 11, 8, 9, 7, 11, 32, 52, 28, 36, 36, 28, 15, 10], 32)
for i, ph in enumerate(d.PHASES):
    r = 5 + i
    wr = wbs_row[ph[1]]
    cell(S_GD, r, 1, ph[0], bold=True, align=CENTER)
    cell(S_GD, r, 2, ph[2], bold=True)
    cell(S_GD, r, 3, f"=WBS!H{wr}", DATE_FMT, color=GREEN, align=CENTER)
    cell(S_GD, r, 4, f"=WBS!I{wr}", DATE_FMT, color=GREEN, align=CENTER)
    cell(S_GD, r, 5, f"=ROUND(NETWORKDAYS(C{r},D{r},NgayNghi)/5,1)", "0.0", align=CENTER)
    cell(S_GD, r, 6, f"=WBS!K{wr}", "0.0#", color=GREEN, align=CENTER)
    g = d.PH_GATE.get(ph[0])
    cell(S_GD, r, 7, g or "—", bold=True, align=CENTER)
    cell(S_GD, r, 8, f"={GATE_DATE(g)}" if g else "", DATE_FMT, color=GREEN, align=CENTER)
    for j, v in enumerate(ph[3:9]):
        cell(S_GD, r, 9 + j, v)
    cell(S_GD, r, 15, f"=WBS!P{wr}", color=GREEN)
    cell(S_GD, r, 16, f"=WBS!Q{wr}", "0%", color=GREEN, align=CENTER)
    S_GD.row_dimensions[r].height = 160
GD_LAST = 4 + len(d.PHASES)
status_cf(S_GD, f"O5:O{GD_LAST}", "O5")
S_GD.freeze_panes = "C5"
page_setup(S_GD, "4:4")

# ---------------------------------------------------------------- Gantt tuần
G_START = date(2026, 9, 21)
WEEKS = (END_BUILD - G_START).days // 7 + 2
GC0 = 11
title(S_G, f"BIỂU ĐỒ GANTT THEO TUẦN (từ 21/09/2026, {WEEKS} tuần đến hết triển khai diện rộng)",
      "Thanh màu vẽ tự động từ ngày Bắt đầu/Kết thúc trong WBS. Ô tuần của gói cấp 3 chứa số ngày công phân bổ trong tuần (ẩn số, dùng cho các dòng tải theo vai trò). Giai đoạn bảo hành xem ở sheet Lộ trình tháng.",
      GC0 + WEEKS - 1,
      "Màu: xanh đậm = giai đoạn · xanh vừa = nhóm công việc · xanh nhạt = gói công việc · xanh lá = hoàn thành · cột hồng = tuần hiện tại · ◆ = mốc/cổng")
glab = ["Mã", "Cấp", "Hạng mục", "Bắt đầu", "Kết thúc", "Công", "Thời lượng", "% HT", "Trạng thái", "Phụ trách"]
for i, (lab, w) in enumerate(zip(glab, [8, 4, 44, 10, 10, 6, 6, 6, 13, 14]), 1):
    c = S_G.cell(5, i, lab)
    c.font = F(True, "FFFFFF")
    c.fill = fill(NAVY)
    c.alignment = CENTER
    c.border = BOX
    S_G.column_dimensions[CL(i)].width = w
for k in range(WEEKS):
    col = GC0 + k
    L, prev = CL(col), CL(col - 1)
    S_G.cell(5, col, "=DATE(2026,9,21)" if k == 0 else f"={prev}5+7").number_format = "dd/mm"
    S_G.cell(4, col, f'="T"&MONTH({L}5)&"/"&RIGHT(YEAR({L}5),2)' if k == 0 else f'=IF(MONTH({L}5)<>MONTH({prev}5),"T"&MONTH({L}5)&"/"&RIGHT(YEAR({L}5),2),"")')
    S_G.cell(6, col, k + 1)
    for rr in (4, 5, 6):
        c = S_G.cell(rr, col)
        c.font = F(rr != 6, "FFFFFF" if rr == 5 else NAVY, 8)
        c.fill = fill(NAVY) if rr == 5 else fill("FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90 if rr == 5 else 0)
        c.border = BOX
    S_G.column_dimensions[L].width = 3.4
S_G.row_dimensions[5].height = 42
S_G.cell(6, 3, "Tuần số").font = F(False, "4A5A70", 8)
GR0 = 7
for i in range(len(rows)):
    r, wr = GR0 + i, FIRST + i
    lvl = {"L1": 1, "L2": 2, "L3": 3}[rows[i][0]]
    bg = {1: L1FILL, 2: L2FILL, 3: None}[lvl]
    cell(S_G, r, 1, f"=WBS!A{wr}", bold=lvl < 3, bg=bg)
    cell(S_G, r, 2, f"=WBS!B{wr}", bg=bg, align=CENTER)
    cell(S_G, r, 3, f"=WBS!C{wr}", bold=lvl < 3, bg=bg).alignment = Alignment(indent=lvl - 1, vertical="center")
    cell(S_G, r, 4, f"=WBS!H{wr}", "dd/mm/yy", bg=bg, align=CENTER)
    cell(S_G, r, 5, f"=WBS!I{wr}", "dd/mm/yy", bg=bg, align=CENTER)
    cell(S_G, r, 6, f"=WBS!K{wr}", "0.#", bg=bg, align=CENTER)
    cell(S_G, r, 7, f"=WBS!J{wr}", "0", bg=bg, align=CENTER)
    cell(S_G, r, 8, f"=WBS!Q{wr}", "0%", bg=bg, align=CENTER)
    cell(S_G, r, 9, f"=WBS!P{wr}", bg=bg)
    cell(S_G, r, 10, f"=WBS!N{wr}", bg=bg)
    for k in range(WEEKS):
        L = CL(GC0 + k)
        c = S_G.cell(r, GC0 + k)
        if lvl == 3:
            c.value = (f'=IF(AND($D{r}<={L}$5+6,$E{r}>={L}$5),IF($H{r}>=1,$F{r},$F{r}*(1-$H{r}))'
                       f'*NETWORKDAYS(MAX($D{r},{L}$5),MIN($E{r},{L}$5+6),NgayNghi)/MAX(NETWORKDAYS($D{r},$E{r},NgayNghi),1),"")')
        c.number_format = ";;;"
        c.border = Border(left=hair, right=hair, top=thin, bottom=thin)
    S_G.row_dimensions[r].outlineLevel = lvl - 1
G_TASK_LAST = GR0 + len(rows) - 1
gr = G_TASK_LAST + 2
cell(S_G, gr, 3, "MỐC VÀ CỔNG", bold=True, bg=L1FILL)
G_GATE0 = gr + 1
gantt_gates = [g for g in d.GATES if g[0] != "G11"]
for i, g in enumerate(gantt_gates):
    r = G_GATE0 + i
    cr = gate_row[g[0]]
    cell(S_G, r, 1, f"='Mốc & Cổng'!A{cr}", bold=True, color="B45309")
    cell(S_G, r, 3, f"='Mốc & Cổng'!B{cr}")
    cell(S_G, r, 4, f"='Mốc & Cổng'!E{cr}", "dd/mm/yy", align=CENTER)
    for k in range(WEEKS):
        L = CL(GC0 + k)
        c = S_G.cell(r, GC0 + k, f'=IF(AND($D{r}>={L}$5,$D{r}<={L}$5+6),"◆","")')
        c.font = F(True, "B45309", 9)
        c.alignment = CENTER
G_GATE_LAST = G_GATE0 + len(gantt_gates) - 1
LOAD0 = G_GATE_LAST + 2
cell(S_G, LOAD0, 3, "TẢI THEO VAI TRÒ", bold=True, bg=L1FILL)
load_row = {k: LOAD0 + 1 + j for j, k in enumerate(ROLE_KEYS)}
TOT_ROW = LOAD0 + 1 + len(ROLE_KEYS)
util_row = {k: TOT_ROW + 1 + j for j, k in enumerate(ROLE_KEYS)}
CAP_ROW = TOT_ROW + 1 + len(ROLE_KEYS)
for k in ROLE_KEYS:
    cell(S_G, load_row[k], 3, f"Ngày công – {ROLE_NAME[k]}", bg=L2FILL)
    cell(S_G, load_row[k], 10, ROLE_NAME[k], bg=L2FILL)
    cell(S_G, util_row[k], 3, f"Mức sử dụng – {ROLE_NAME[k]}", bg=L2FILL)
    cell(S_G, util_row[k], 10, ROLE_NAME[k], bg=L2FILL)
cell(S_G, TOT_ROW, 3, "Ngày công – cả đội", bold=True, bg=L2FILL)
cell(S_G, CAP_ROW, 3, "Năng lực mỗi người (ngày làm việc trong tuần)", bg=L2FILL)
for k_ in range(WEEKS):
    col = GC0 + k_
    L = CL(col)
    cells_ = [(S_G.cell(CAP_ROW, col, f"=NETWORKDAYS({L}5,{L}5+4,NgayNghi)"), "0", False)]
    for k in ROLE_KEYS:
        cells_.append((S_G.cell(load_row[k], col, f"=SUMIF($J${GR0}:$J${G_TASK_LAST},$J{load_row[k]},{L}{GR0}:{L}{G_TASK_LAST})"), "0.0", False))
        cells_.append((S_G.cell(util_row[k], col, f"=IFERROR({L}{load_row[k]}/{L}{CAP_ROW},0)"), "0%", False))
    cells_.append((S_G.cell(TOT_ROW, col, f"=SUM({L}{load_row[ROLE_KEYS[0]]}:{L}{load_row[ROLE_KEYS[-1]]})"), "0.0", True))
    for x_, fmt, b_ in cells_:
        x_.number_format = fmt
        x_.font = F(b_, INK, 8)
        x_.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
        x_.border = BOX
for rr in list(load_row.values()) + list(util_row.values()) + [TOT_ROW, CAP_ROW]:
    S_G.row_dimensions[rr].height = 28
fc = CL(GC0)
grid = f"{fc}{GR0}:{CL(GC0 + WEEKS - 1)}{G_TASK_LAST}"
cond = f"$D{GR0}<={fc}$5+6,$E{GR0}>={fc}$5"
S_G.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($I{GR0}="Hoàn thành",{cond})'], fill=fill("7CC49A"), stopIfTrue=True))
S_G.conditional_formatting.add(grid, FormulaRule(formula=[f"AND($B{GR0}=1,{cond})"], fill=fill(NAVY), stopIfTrue=True))
S_G.conditional_formatting.add(grid, FormulaRule(formula=[f"AND($B{GR0}=2,{cond})"], fill=fill(NAVY2), stopIfTrue=True))
S_G.conditional_formatting.add(grid, FormulaRule(formula=[f"AND($B{GR0}=3,{cond})"], fill=fill(NAVY3), stopIfTrue=True))
S_G.conditional_formatting.add(f"{fc}{GR0}:{CL(GC0 + WEEKS - 1)}{G_GATE_LAST}", FormulaRule(formula=[f"AND(TODAY()>={fc}$5,TODAY()<={fc}$5+6)"], fill=fill("FBE7EA")))
u0, u1 = min(util_row.values()), max(util_row.values())
S_G.conditional_formatting.add(f"{fc}{u0}:{CL(GC0 + WEEKS - 1)}{u1}", FormulaRule(formula=[f"{fc}{u0}>1.05"], fill=fill("FBE9E9"), font=Font(name=FONT, color="B91C1C", bold=True)))
S_G.sheet_properties.outlinePr.summaryBelow = False
S_G.freeze_panes = S_G.cell(GR0, GC0)
for col in ("B", "G", "H"):
    S_G.column_dimensions[col].hidden = True
page_setup(S_G, "4:6")

# ---------------------------------------------------------------- Lộ trình tháng
M_START = date(2026, 9, 1)
MONTHS = (END_ALL.year - M_START.year) * 12 + END_ALL.month - M_START.month + 1
MC0 = 6
title(S_M, f"LỘ TRÌNH TỔNG THỂ THEO THÁNG (09/2026 – {END_ALL.month:02d}/{END_ALL.year})",
      "Giai đoạn (cấp 1) và nhóm công việc (cấp 2), tự cập nhật theo WBS. ◆ = mốc/cổng trong tháng.", MC0 + MONTHS - 1,
      "Màu: xanh đậm = giai đoạn · xanh vừa = nhóm công việc · xanh lá = hoàn thành · cột hồng = tháng hiện tại")
for i, (lab, w) in enumerate(zip(["Mã", "Cấp", "Hạng mục", "Bắt đầu", "Kết thúc"], [8, 4, 46, 11, 11]), 1):
    c = S_M.cell(5, i, lab)
    c.font = F(True, "FFFFFF")
    c.fill = fill(NAVY)
    c.alignment = CENTER
    c.border = BOX
    S_M.column_dimensions[CL(i)].width = w
for k in range(MONTHS):
    col = MC0 + k
    L, prev = CL(col), CL(col - 1)
    S_M.cell(5, col, "=DATE(2026,9,1)" if k == 0 else f"=EDATE({prev}5,1)").number_format = "mm/yy"
    S_M.cell(4, col, f"=YEAR({L}5)" if k == 0 else f'=IF(YEAR({L}5)<>YEAR({prev}5),YEAR({L}5),"")')
    for rr in (4, 5):
        c = S_M.cell(rr, col)
        c.font = F(True, "FFFFFF" if rr == 5 else NAVY, 9)
        c.fill = fill(NAVY) if rr == 5 else fill("FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90 if rr == 5 else 0)
        c.border = BOX
    S_M.column_dimensions[L].width = 3.6
S_M.row_dimensions[5].height = 36
r = 6
m_first = r
for i, (kind, it, gd) in enumerate(rows):
    if kind == "L3":
        continue
    wr = FIRST + i
    lvl = 1 if kind == "L1" else 2
    bg = L1FILL if lvl == 1 else None
    cell(S_M, r, 1, f"=WBS!A{wr}", bold=lvl == 1, bg=bg)
    cell(S_M, r, 2, f"=WBS!B{wr}", bg=bg, align=CENTER)
    cell(S_M, r, 3, f"=WBS!C{wr}", bold=lvl == 1, bg=bg).alignment = Alignment(indent=lvl - 1, vertical="center")
    cell(S_M, r, 4, f"=WBS!H{wr}", DATE_FMT, bg=bg, align=CENTER)
    cell(S_M, r, 5, f"=WBS!I{wr}", DATE_FMT, bg=bg, align=CENTER)
    for k in range(MONTHS):
        S_M.cell(r, MC0 + k).border = BOX
    S_M.row_dimensions[r].outlineLevel = lvl - 1
    r += 1
m_last = r - 1
r += 1
cell(S_M, r, 3, "MỐC VÀ CỔNG", bold=True, bg=L1FILL)
r += 1
for g in d.GATES:
    cr = gate_row[g[0]]
    cell(S_M, r, 1, f"='Mốc & Cổng'!A{cr}", bold=True, color="B45309")
    cell(S_M, r, 3, f"='Mốc & Cổng'!B{cr}")
    cell(S_M, r, 4, f"='Mốc & Cổng'!E{cr}", DATE_FMT, align=CENTER)
    for k in range(MONTHS):
        L = CL(MC0 + k)
        c = S_M.cell(r, MC0 + k, f'=IF(AND($D{r}>={L}$5,$D{r}<=EOMONTH({L}$5,0)),"◆","")')
        c.font = F(True, "B45309", 10)
        c.alignment = CENTER
        c.border = BOX
    r += 1
mg_last = r - 1
mf = CL(MC0)
mgrid = f"{mf}{m_first}:{CL(MC0 + MONTHS - 1)}{m_last}"
mcond = f"$D{m_first}<=EOMONTH({mf}$5,0),$E{m_first}>={mf}$5"
S_M.conditional_formatting.add(mgrid, FormulaRule(formula=[f'AND(INDEX(WBS!$P:$P,MATCH($A{m_first},WBS!$A:$A,0))="Hoàn thành",{mcond})'], fill=fill("7CC49A"), stopIfTrue=True))
S_M.conditional_formatting.add(mgrid, FormulaRule(formula=[f"AND($B{m_first}=1,{mcond})"], fill=fill(NAVY), stopIfTrue=True))
S_M.conditional_formatting.add(mgrid, FormulaRule(formula=[f"AND($B{m_first}=2,{mcond})"], fill=fill(NAVY2), stopIfTrue=True))
S_M.conditional_formatting.add(f"{mf}{m_first}:{CL(MC0 + MONTHS - 1)}{mg_last}", FormulaRule(formula=[f"AND(TODAY()>={mf}$5,TODAY()<=EOMONTH({mf}$5,0))"], fill=fill("FBE7EA")))
S_M.column_dimensions["B"].hidden = True
S_M.sheet_properties.outlinePr.summaryBelow = False
S_M.freeze_panes = S_M.cell(6, MC0)
page_setup(S_M, "4:5")

# ---------------------------------------------------------------- Phạm vi tính năng
PCOLS = ["Mã", "Epic", "Tên epic", "Tính năng", "Phát hành", "Ưu tiên", "Nền tảng", "Mã WBS", "Hạng mục WBS thực hiện", "Hoàn thành dự kiến", "Trạng thái"]
title(S_P, "PHẠM VI TÍNH NĂNG THEO PHÁT HÀNH",
      "Nguồn: docs/02-requirements/feature-map.md. Nền tảng: W = web (MVP), D = ứng dụng Windows, M = ứng dụng Android. Hạng mục, ngày và trạng thái tra từ WBS theo Mã WBS.", 18)
header(S_P, 4, PCOLS, [7, 6, 26, 52, 11, 8, 12, 9, 44, 12, 15])
for i, f_ in enumerate(d.FEATURES):
    r = 5 + i
    cell(S_P, r, 1, f_[0], bold=True)
    cell(S_P, r, 2, f_[1], align=CENTER)
    cell(S_P, r, 3, d.EPICS[f_[1]])
    cell(S_P, r, 4, f_[2])
    cell(S_P, r, 5, f_[3], color=BLUE, align=CENTER)
    cell(S_P, r, 6, f_[4], color=BLUE, align=CENTER)
    cell(S_P, r, 7, f_[5], align=CENTER)
    cell(S_P, r, 8, f_[6], color=BLUE, align=CENTER)
    cell(S_P, r, 9, f'=IFERROR(INDEX(WBS!$C:$C,MATCH($H{r},WBS!$A:$A,0)),"—")')
    cell(S_P, r, 10, f'=IFERROR(INDEX(WBS!$I:$I,MATCH($H{r},WBS!$A:$A,0)),"")', DATE_FMT, align=CENTER)
    cell(S_P, r, 11, f'=IFERROR(INDEX(WBS!$P:$P,MATCH($H{r},WBS!$A:$A,0)),"")')
P_LAST = 4 + len(d.FEATURES)
dvp = DataValidation(type="list", formula1="=DM_PhatHanh", allow_blank=True)
dvp.add(f"E5:E{P_LAST}")
dvu = DataValidation(type="list", formula1="=DM_UuTien", allow_blank=True)
dvu.add(f"F5:F{P_LAST}")
S_P.add_data_validation(dvp)
S_P.add_data_validation(dvu)
S_P.auto_filter.ref = f"A4:K{P_LAST}"
for val, bg_ in [("MVP", "E6ECF8"), ("Ứng dụng", "E3F1F9"), ("R2", "FDF1E3"), ("R3", "F3E8FF"), ("Thí điểm", "E7F5EC"), ("Triển khai", "E7F5EC"), ("Bảo hành", "EEF2F7")]:
    S_P.conditional_formatting.add(f"E5:E{P_LAST}", FormulaRule(formula=[f'E5="{val}"'], fill=fill(bg_)))
S_P.column_dimensions["L"].width = 3
sc = 13
cell(S_P, 4, sc, "Phát hành", bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
S_P.column_dimensions[CL(sc)].width = 13
for j, p in enumerate(["P0", "P1", "P2", "Tổng"]):
    cell(S_P, 4, sc + 1 + j, p, bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
    S_P.column_dimensions[CL(sc + 1 + j)].width = 8
for i, rel in enumerate(d.RELEASES + ["Tổng"]):
    r = 5 + i
    tot_ = rel == "Tổng"
    bgx = L2FILL if tot_ else None
    cell(S_P, r, sc, rel, bold=tot_, bg=bgx)
    for j, p in enumerate(["P0", "P1", "P2"]):
        colL = CL(sc + 1 + j)
        v = f"=SUM({colL}5:{colL}{r - 1})" if tot_ else f'=COUNTIFS($E$5:$E${P_LAST},${CL(sc)}{r},$F$5:$F${P_LAST},"{p}")'
        cell(S_P, r, sc + 1 + j, v, "0", bold=tot_, bg=bgx, align=CENTER)
    cell(S_P, r, sc + 4, f"=SUM({CL(sc + 1)}{r}:{CL(sc + 3)}{r})", "0", bold=True, bg=bgx, align=CENTER)
S_P.freeze_panes = "E5"
page_setup(S_P, "4:4")

# ---------------------------------------------------------------- Nguồn lực: analysis
Wr = lambda col: f"WBS!${col}${FIRST}:${col}${LAST}"
CMP = PR + 3
cell(S_N, CMP, 2, "So sánh với mốc MVP mục tiêu", bold=True, bg=L2FILL)
cell(S_N, CMP, 4, "", bg=L2FILL)
cmp_rows = [
    ("Ngày nghiệm thu MVP theo kế hoạch (cổng G5)", f"={GATE_DATE('G5')}", DATE_FMT),
    ("Số ngày làm việc còn dư (+) hoặc trễ (−) so với mục tiêu", f"=IF(D{CMP + 1}<={TARGET},NETWORKDAYS(D{CMP + 1},{TARGET},NgayNghi)-1,-(NETWORKDAYS({TARGET},D{CMP + 1},NgayNghi)-1))", "0"),
    ("Đánh giá", f'=IF(D{CMP + 1}<={TARGET},"Đạt mục tiêu","Trễ mục tiêu")', "@"),
    ("Công đến nghiệm thu MVP (GĐ0–GĐ3, ngày)", f'=SUM(SUMIFS({Wr("K")},{Wr("E")},{{"GĐ0","GĐ1","GĐ2","GĐ3"}},{Wr("B")},3))', "0.0"),
]
for i, (lab, fx, fmt) in enumerate(cmp_rows):
    cell(S_N, CMP + 1 + i, 2, lab)
    cell(S_N, CMP + 1 + i, 4, fx, fmt, bold=True, align=CENTER)
S_N.conditional_formatting.add(f"D{CMP + 3}", FormulaRule(formula=[f'D{CMP + 3}="Đạt mục tiêu"'], fill=fill("E7F5EC"), font=Font(name=FONT, color="15803D", bold=True)))
S_N.conditional_formatting.add(f"D{CMP + 3}", FormulaRule(formula=[f'D{CMP + 3}="Trễ mục tiêu"'], fill=fill("FBE9E9"), font=Font(name=FONT, color="B91C1C", bold=True)))
NT = CMP + 7
cell(S_N, NT - 1, 1, "Khối lượng theo giai đoạn và vai trò (ngày công)", bold=True, border=False)
ncols = ["Mã GĐ", "Giai đoạn", "Bắt đầu", "Kết thúc", "Số tuần"] + [ROLE_NAME[k] for k in ROLE_KEYS] + ["Tổng", "Mức sử dụng đội"]
header(S_N, NT, ncols, None, 30)
for i, w in enumerate([8, 50, 12, 12, 9, 13, 13, 13, 10, 12], 1):
    S_N.column_dimensions[CL(i)].width = w
for i, ph in enumerate(d.PHASES):
    r = NT + 1 + i
    g_ = 5 + i
    cell(S_N, r, 1, ph[0], bold=True, align=CENTER)
    cell(S_N, r, 2, ph[2])
    cell(S_N, r, 3, f"='Giai đoạn & Phạm vi'!C{g_}", DATE_FMT, color=GREEN, align=CENTER)
    cell(S_N, r, 4, f"='Giai đoạn & Phạm vi'!D{g_}", DATE_FMT, color=GREEN, align=CENTER)
    cell(S_N, r, 5, f"='Giai đoạn & Phạm vi'!E{g_}", "0.0", color=GREEN, align=CENTER)
    for j, k in enumerate(ROLE_KEYS):
        cell(S_N, r, 6 + j, f'=SUMIFS({Wr("K")},{Wr("E")},$A{r},{Wr("B")},3,{Wr("N")},"{ROLE_NAME[k]}")', "0.0#", align=CENTER)
    tc = 6 + len(ROLE_KEYS)
    cell(S_N, r, tc, f"=SUM({CL(6)}{r}:{CL(tc - 1)}{r})", "0.0#", bold=True, align=CENTER)
    cell(S_N, r, tc + 1, f"=IFERROR({CL(tc)}{r}/(E{r}*5*{len(ROLE_KEYS)}),0)", "0%", align=CENTER)
N_LAST = NT + len(d.PHASES)
rt = N_LAST + 1
cell(S_N, rt, 2, "Tổng", bold=True, bg=L2FILL)
for c_ in (1, 3, 4, 5, 6 + len(ROLE_KEYS) + 1):
    cell(S_N, rt, c_, "", bg=L2FILL)
for j in range(len(ROLE_KEYS) + 1):
    colL = CL(6 + j)
    cell(S_N, rt, 6 + j, f"=SUM({colL}{NT + 1}:{colL}{N_LAST})", "0.0#", True, bg=L2FILL, align=CENTER)
kr = rt + 2
cell(S_N, kr, 1, "Mức sử dụng theo tuần (từ sheet Gantt tuần)", bold=True, border=False)
header(S_N, kr + 1, ["Mã", "Vai trò", "Cao nhất", "Trung bình đến nghiệm thu MVP", "Số tuần trên 105%"], None, 30)
mvp_weeks = (group_end("3") - G_START).days // 7 + 1
for j, k in enumerate(ROLE_KEYS):
    r = kr + 2 + j
    rowref = f"'Gantt tuần'!{CL(GC0)}{util_row[k]}:{CL(GC0 + WEEKS - 1)}{util_row[k]}"
    mvpref = f"'Gantt tuần'!{CL(GC0 + 1)}{util_row[k]}:{CL(GC0 + mvp_weeks - 1)}{util_row[k]}"
    cell(S_N, r, 1, k, align=CENTER)
    cell(S_N, r, 2, ROLE_NAME[k])
    cell(S_N, r, 3, f"=MAX({rowref})", "0%", bold=True, align=CENTER)
    cell(S_N, r, 4, f"=AVERAGE({mvpref})", "0%", bold=True, align=CENTER)
    cell(S_N, r, 5, f'=COUNTIF({rowref},">1.05")', "0", bold=True, align=CENTER)
tbl0 = kr + 3 + len(ROLE_KEYS) + 2
cell(S_N, tbl0 - 1, 1, "Dữ liệu biểu đồ tải tuần", bold=True, border=False)
labs = ["Tuần", "Bắt đầu tuần"] + [ROLE_NAME[k] for k in ROLE_KEYS] + ["Năng lực đội"]
for i, lab in enumerate(labs, 1):
    cell(S_N, tbl0, i, lab, bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
for k_ in range(WEEKS):
    r = tbl0 + 1 + k_
    L = CL(GC0 + k_)
    cell(S_N, r, 1, k_ + 1, "0", align=CENTER)
    cell(S_N, r, 2, f"='Gantt tuần'!{L}5", "dd/mm/yy", color=GREEN, align=CENTER)
    for j, k in enumerate(ROLE_KEYS):
        cell(S_N, r, 3 + j, f"='Gantt tuần'!{L}{load_row[k]}", "0.0", color=GREEN, align=CENTER)
    cell(S_N, r, 3 + len(ROLE_KEYS), f"='Gantt tuần'!{L}{CAP_ROW}*{len(ROLE_KEYS)}", "0", color=GREEN, align=CENTER)
tbl_last = tbl0 + WEEKS
bar = BarChart()
bar.type = "col"
bar.grouping = "stacked"
bar.overlap = 100
bar.title = "Ngày công theo tuần và vai trò so với năng lực đội"
bar.y_axis.title = "Ngày công"
bar.x_axis.title = "Tuần bắt đầu"
bar.add_data(Reference(S_N, min_col=3, max_col=2 + len(ROLE_KEYS), min_row=tbl0, max_row=tbl_last), titles_from_data=True)
bar.set_categories(Reference(S_N, min_col=2, min_row=tbl0 + 1, max_row=tbl_last))
for s_, colr in zip(bar.series, [NAVY, NAVY2, "7CC49A"]):
    s_.graphicalProperties.solidFill = colr
line = LineChart()
line.add_data(Reference(S_N, min_col=3 + len(ROLE_KEYS), min_row=tbl0, max_row=tbl_last), titles_from_data=True)
line.series[0].graphicalProperties.line.solidFill = "B91C1C"
line.series[0].graphicalProperties.line.dashStyle = "dash"
bar += line
bar.height, bar.width = 10, 32
bar.legend.position = "b"
S_N.add_chart(bar, "L4")
page_setup(S_N, None)

# ---------------------------------------------------------------- Rủi ro
RCOLS = ["Mã", "Rủi ro", "Giai đoạn ảnh hưởng", "Khả năng (1–3)", "Tác động (1–3)", "Điểm", "Mức", "Biện pháp giảm thiểu", "Dấu hiệu cảnh báo", "Phụ trách", "Trạng thái"]
title(S_R, "SỔ RỦI RO", "Điểm = Khả năng × Tác động. Mức: Cao ≥ 6, Trung bình 3–5, Thấp ≤ 2. Rà soát hàng tuần trong báo cáo tiến độ.", 11)
header(S_R, 4, RCOLS, [6, 48, 13, 10, 10, 7, 12, 52, 36, 14, 15])
for i, rk in enumerate(d.RISKS):
    r = 5 + i
    cell(S_R, r, 1, rk[0], bold=True, align=CENTER)
    cell(S_R, r, 2, rk[1])
    cell(S_R, r, 3, rk[2], align=CENTER)
    cell(S_R, r, 4, rk[3], "0", color=BLUE, align=CENTER)
    cell(S_R, r, 5, rk[4], "0", color=BLUE, align=CENTER)
    cell(S_R, r, 6, f"=D{r}*E{r}", "0", bold=True, align=CENTER)
    cell(S_R, r, 7, f'=IF(F{r}>=6,"Cao",IF(F{r}>=3,"Trung bình","Thấp"))', bold=True, align=CENTER)
    cell(S_R, r, 8, rk[5])
    cell(S_R, r, 9, rk[6])
    cell(S_R, r, 10, d.TDA, color=BLUE)
    cell(S_R, r, 11, "Đang theo dõi", color=BLUE)
R_LAST = 4 + len(d.RISKS)
for val, bg_, fg_ in [("Cao", "FBE9E9", "B91C1C"), ("Trung bình", "FDF1E3", "B45309"), ("Thấp", "E7F5EC", "15803D")]:
    S_R.conditional_formatting.add(f"G5:G{R_LAST}", FormulaRule(formula=[f'G5="{val}"'], fill=fill(bg_), font=Font(name=FONT, color=fg_, bold=True)))
dvr = DataValidation(type="list", formula1="=DM_RuiRo", allow_blank=True)
dvr.add(f"K5:K{R_LAST}")
dvn = DataValidation(type="whole", operator="between", formula1="1", formula2="3", error="Nhập 1, 2 hoặc 3", errorTitle="Thang điểm")
dvn.add(f"D5:E{R_LAST}")
S_R.add_data_validation(dvr)
S_R.add_data_validation(dvn)
S_R.auto_filter.ref = f"A4:K{R_LAST}"
mr = R_LAST + 3
cell(S_R, mr - 1, 2, "Ma trận rủi ro (số rủi ro chưa đóng theo Khả năng × Tác động)", bold=True, border=False)
cell(S_R, mr, 2, "Khả năng \\ Tác động", bold=True, bg=L2FILL, align=CENTER)
for j, imp in enumerate([1, 2, 3]):
    cell(S_R, mr, 3 + j, imp, "0", bold=True, bg=L2FILL, align=CENTER)
for i, lk in enumerate([3, 2, 1]):
    r = mr + 1 + i
    cell(S_R, r, 2, lk, "0", bold=True, bg=L2FILL, align=CENTER)
    for j, imp in enumerate([1, 2, 3]):
        score = lk * imp
        bg_ = "FBE9E9" if score >= 6 else ("FDF1E3" if score >= 3 else "E7F5EC")
        cell(S_R, r, 3 + j, f'=COUNTIFS($D$5:$D${R_LAST},{lk},$E$5:$E${R_LAST},{imp},$K$5:$K${R_LAST},"<>Đã đóng")', "0", True, bg=bg_, align=CENTER)
S_R.freeze_panes = "C5"
page_setup(S_R, "4:4")

# ---------------------------------------------------------------- Giả định
title(S_A, "GIẢ ĐỊNH, PHỤ THUỘC VÀ RÀNG BUỘC", "Mọi con số ước lượng trong workbook dựa trên các giả định dưới đây. Khi một giả định thay đổi, cập nhật WBS và ghi vào nhật ký thay đổi.", 6)
header(S_A, 4, ["Mã", "Loại", "Nội dung", "Ảnh hưởng đến", "Cách xác nhận / Nguồn", "Trạng thái"], [6, 12, 74, 16, 30, 15])
for i, a in enumerate(d.ASSUMPTIONS):
    r = 5 + i
    cell(S_A, r, 1, a[0], bold=True, align=CENTER)
    for j, v in enumerate(a[1:]):
        cell(S_A, r, 2 + j, v, color=BLUE if j == 4 else INK)
A_LAST = 4 + len(d.ASSUMPTIONS)
S_A.conditional_formatting.add(f"F5:F{A_LAST}", FormulaRule(formula=['F5="Đã xác nhận"'], fill=fill("E7F5EC"), font=Font(name=FONT, color="15803D", bold=True)))
S_A.conditional_formatting.add(f"F5:F{A_LAST}", FormulaRule(formula=['F5="Chờ xác nhận"'], fill=fill("FDF1E3"), font=Font(name=FONT, color="B45309", bold=True)))
S_A.freeze_panes = "A5"
page_setup(S_A, "4:4")

# ---------------------------------------------------------------- Tổng quan
ws = S_TQ
for col, w in zip("ABCDEFGH", [3, 48, 22, 14, 14, 12, 12, 46]):
    ws.column_dimensions[col].width = w
ws.merge_cells("B1:H1")
ws["B1"] = "KẾ HOẠCH TỔNG THỂ DỰ ÁN AI MINUTE"
ws["B1"].font = F(True, NAVY, 18)
ws.row_dimensions[1].height = 30
ws.merge_cells("B2:H2")
ws["B2"] = "Thư ký cuộc họp AI bảo mật cho cơ quan nhà nước · WBS, tiến độ và phạm vi từng giai đoạn cho toàn bộ vòng đời dự án"
ws["B2"].font = F(False, "4A5A70", 10, True)
info = [
    ("Tên dự án", "AI Minute – Thư ký cuộc họp AI bảo mật"),
    ("Khách hàng", "Cơ quan nhà nước cấp tỉnh, xã, phường"),
    ("Mục tiêu", "Thư ký mở ứng dụng trong phòng họp, bấm ghi âm, xem bản ghi tiếng Việt theo thời gian thực và nhận biên bản nháp theo mẫu ngay sau cuộc họp. Toàn bộ xử lý chạy trên máy chủ nội bộ, không qua internet."),
    ("Nền tảng theo giai đoạn", "MVP: bản web (ghi âm bằng laptop với Chrome/Edge; điện thoại để xem và duyệt biên bản). GĐ5: ứng dụng Windows và Android – cách dùng được khuyến nghị của sản phẩm cuối (ghi âm khi chạy nền, khóa màn hình, mất mạng). Không làm iOS/macOS."),
    ("Nguồn lực", "Đội 3 người toàn thời gian: trưởng dự án, lập trình viên fullstack, kiểm thử viên (chi tiết ở sheet Nguồn lực). Bảo hành: trưởng dự án."),
    ("Cơ sở ước lượng", "Ngày công ước lượng cho từng gói, đã gồm thiết kế chi tiết, kiểm thử và tài liệu. Mỗi vai trò làm lần lượt các gói của mình; gói bắt đầu khi các gói phụ thuộc đã xong."),
    ("Phiên bản / Ngày lập", "2.0 · 25/09/2026"),
    ("Người lập", "Trưởng dự án"),
]
cell(ws, 4, 2, "THÔNG TIN DỰ ÁN", bold=True, color="FFFFFF", bg=NAVY)
ws.merge_cells("C4:H4")
cell(ws, 4, 3, "", bg=NAVY)
for i, (k, v) in enumerate(info):
    r = 5 + i
    cell(ws, r, 2, k, bold=True, bg=L2FILL, align=VC)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    cell(ws, r, 3, v, align=VC)
    ws.row_dimensions[r].height = 44 if len(v) > 150 else (32 if len(v) > 90 else 18)
r0 = 5 + len(info) + 1
cell(ws, r0, 2, "SỐ LIỆU CHÍNH (tự tính)", bold=True, color="FFFFFF", bg=NAVY)
ws.merge_cells(start_row=r0, start_column=3, end_row=r0, end_column=8)
cell(ws, r0, 3, "", bg=NAVY)
key = [
    ("Bắt đầu dự án", f"=WBS!H{W_TOT}", DATE_FMT),
    ("Nghiệm thu MVP web (cổng G5)", f"={GATE_DATE('G5')}", DATE_FMT),
    ("So với mục tiêu trước 12/2026", f"='Nguồn lực'!D{CMP + 3}&\" (\"&TEXT('Nguồn lực'!D{CMP + 2},\"0\")&\" ngày LV)\"", "@"),
    ("Phát hành ứng dụng Windows và Android (cổng G7)", f"={GATE_DATE('G7')}", DATE_FMT),
    ("Nghiệm thu R2 (cổng G8)", f"={GATE_DATE('G8')}", DATE_FMT),
    ("Nghiệm thu R3 (cổng G9)", f"={GATE_DATE('G9')}", DATE_FMT),
    ("Nghiệm thu tổng thể và chuyển giao (cổng G10)", f"={GATE_DATE('G10')}", DATE_FMT),
    ("Kết thúc bảo hành (cổng G11)", f"={GATE_DATE('G11')}", DATE_FMT),
    ("Số gói công việc (cấp 3)", f"=COUNTIF({Wr('B')},3)", "0"),
    ("Tổng công toàn dự án (ngày)", f"=WBS!K{W_TOT}", "0.0"),
    ("Công đến nghiệm thu MVP (ngày)", f"='Nguồn lực'!D{CMP + 4}", "0.0"),
    ("% hoàn thành toàn dự án (theo công)", f"=WBS!Q{W_TOT}", "0.0%"),
    ("Số cổng đã đạt / tổng số cổng", f"=COUNTIF('Mốc & Cổng'!I5:I{C_LAST},\"Đạt\")&\" / \"&COUNTA('Mốc & Cổng'!A5:A{C_LAST})", "@"),
    ("Số rủi ro mức Cao đang theo dõi", f"=COUNTIFS('Rủi ro'!G5:G{R_LAST},\"Cao\",'Rủi ro'!K5:K{R_LAST},\"<>Đã đóng\")", "0"),
]
for i, (k, fx, fmt) in enumerate(key):
    r = r0 + 1 + i
    cell(ws, r, 2, k, bold=True, bg=L2FILL, align=VC)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    cell(ws, r, 3, fx, fmt, bold=True, color=GREEN, align=Alignment(horizontal="left", vertical="center"))
    ws.cell(r, 4).border = BOX
p0 = r0 + len(key) + 2
cell(ws, p0, 2, "TÓM TẮT GIAI ĐOẠN", bold=True, color="FFFFFF", bg=NAVY)
for j, lab in enumerate(["Mã GĐ", "Bắt đầu", "Kết thúc", "Số tuần", "Công (ngày)", "Tiêu chí kết thúc (cổng)"]):
    cell(ws, p0, 3 + j, lab, bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
for i, ph in enumerate(d.PHASES):
    r = p0 + 1 + i
    g_ = 5 + i
    cell(ws, r, 2, f"='Giai đoạn & Phạm vi'!B{g_}", bold=True, color=GREEN, align=VC)
    cell(ws, r, 3, f"='Giai đoạn & Phạm vi'!A{g_}", color=GREEN, align=CENTER)
    cell(ws, r, 4, f"='Giai đoạn & Phạm vi'!C{g_}", DATE_FMT, color=GREEN, align=CENTER)
    cell(ws, r, 5, f"='Giai đoạn & Phạm vi'!D{g_}", DATE_FMT, color=GREEN, align=CENTER)
    cell(ws, r, 6, f"='Giai đoạn & Phạm vi'!E{g_}", "0.0", color=GREEN, align=CENTER)
    cell(ws, r, 7, f"='Giai đoạn & Phạm vi'!F{g_}", "0.0#", color=GREEN, align=CENTER)
    cell(ws, r, 8, f"='Giai đoạn & Phạm vi'!M{g_}", color=GREEN, align=VC)
    ws.row_dimensions[r].height = 58
s0 = p0 + len(d.PHASES) + 2
cell(ws, s0, 2, "HƯỚNG DẪN ĐỌC", bold=True, color="FFFFFF", bg=NAVY)
ws.merge_cells(start_row=s0, start_column=3, end_row=s0, end_column=8)
cell(ws, s0, 3, "", bg=NAVY)
guide = [
    ("Giai đoạn & Phạm vi", "Mục tiêu, phạm vi, ngoài phạm vi, sản phẩm bàn giao, cổng và tiêu chí kết thúc của 10 giai đoạn và phần quản lý dự án."),
    ("WBS", "Toàn bộ gói công việc theo 3 cấp, người phụ trách, phụ thuộc và tiến độ tự lập. Chỉ nhập liệu ở dòng cấp 3 (chữ xanh dương)."),
    ("Gantt tuần", "Tiến độ theo tuần đến hết triển khai diện rộng, kèm tải và mức sử dụng của từng vai trò mỗi tuần."),
    ("Lộ trình tháng", "Toàn cảnh theo tháng đến hết bảo hành, gồm các cổng."),
    ("Mốc & Cổng", "11 cổng kiểm soát; ngày tự tính từ WBS; cập nhật kết quả khi họp cổng."),
    ("Phạm vi tính năng", "Tính năng theo phát hành (MVP, Thí điểm, Ứng dụng, R2, R3, Triển khai, Bảo hành) và ưu tiên, nối với WBS."),
    ("Nguồn lực", "Vai trò, tỷ lệ thời gian (nền vàng), so sánh với mốc MVP mục tiêu, khối lượng theo giai đoạn và vai trò, tải theo tuần."),
    ("Rủi ro", f"{len(d.RISKS)} rủi ro, điểm và mức tự tính, ma trận rủi ro."),
    ("Giả định & Ràng buộc", "Cơ sở của mọi ước lượng; cần xác nhận các dòng 'Chờ xác nhận'."),
    ("Ngày nghỉ, Danh mục", "Ngày nghỉ lễ dùng để lập lịch; giá trị cho danh sách chọn và ý nghĩa cách lập lịch."),
    ("Quy ước màu", "Chữ xanh dương = số liệu nhập · chữ đen = công thức · chữ xanh lá = lấy từ sheet khác · nền vàng = tham số."),
]
for i, (k, v) in enumerate(guide):
    r = s0 + 1 + i
    cell(ws, r, 2, k, bold=True, bg=L2FILL, align=VC)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    cell(ws, r, 3, v, align=VC)
v0 = s0 + len(guide) + 2
cell(ws, v0, 2, "LỊCH SỬ PHIÊN BẢN", bold=True, color="FFFFFF", bg=NAVY)
for j, lab in enumerate(["Phiên bản", "Ngày", "Người lập"]):
    cell(ws, v0, 3 + j, lab, bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
ws.merge_cells(start_row=v0, start_column=6, end_row=v0, end_column=8)
cell(ws, v0, 6, "Nội dung thay đổi", bold=True, color="FFFFFF", bg=NAVY)
hist = [("1.0", date(2026, 9, 25), "Lập kế hoạch tổng thể lần đầu sau khi chốt phạm vi mới với khách hàng (24/09/2026)."),
        ("1.1", date(2026, 9, 25), "Ước lượng lại; lập lịch tự động theo khối lượng; các mốc tính lại từ WBS."),
        ("2.0", date(2026, 9, 25), "MVP làm bản web trước, không làm kiểm chứng riêng; ứng dụng Windows và Android thành giai đoạn riêng (GĐ5); đội 3 người; lập lịch theo vai trò và phụ thuộc; mục tiêu nghiệm thu MVP trước 12/2026.")]
for i, (ver, dt, txt) in enumerate(hist):
    r = v0 + 1 + i
    cell(ws, r, 2, "")
    cell(ws, r, 3, ver, align=CENTER)
    cell(ws, r, 4, dt, DATE_FMT, align=CENTER)
    cell(ws, r, 5, "Trưởng dự án", align=CENTER)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
    cell(ws, r, 6, txt)
    ws.row_dimensions[r].height = 44 if len(txt) > 120 else 30
ws.sheet_view.showGridLines = False
page_setup(ws, None, portrait=True)

for s_ in wb.worksheets:
    s_.sheet_view.zoomScale = 90
S_TQ.sheet_view.zoomScale = 100
wb.save(OUT)

print("saved", OUT, "| weeks", WEEKS, "months", MONTHS)
for g, code in d.GATE_WBS.items():
    print(g, group_end(code))
for k in ROLE_KEYS:
    print(k, "packages", len(chains[k]), "effort", round(sum(BYCODE[c][8] for c in chains[k]), 2))
print("MVP effort", sum(it[8] for it in L3ROWS if it[1].split(".")[0] in ("0", "1", "2", "3")))
