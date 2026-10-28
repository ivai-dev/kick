"""Baseline estimate workbook: WBS, effort and schedule as live Excel formulas. It tracks nothing.

Inputs (blue): plan start, each package's effort (or effort rate), fixed dates, role focus, hard dates, holidays.
Everything else is a formula with the same rules as kick.schedule.build_schedule(..., baseline=True).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .labels import label
from .model import Plan, Project, code_key, text
from .schedule import role_queues

HEADER = 4  # header row of every table sheet; rows 1-2 hold the title and the note
FIRST = 5  # first data row
SPARE_HOLIDAYS = 50  # empty rows inside HolidayDates for dates added by hand
WBS_COLUMNS = ("code", "level", "name", "desc", "phase", "release", "prio", "role", "mode", "deps", "follows",
               "effort", "rate", "days", "start", "end", "ref", "coop", "note")
COL = {key: get_column_letter(i + 1) for i, key in enumerate(WBS_COLUMNS)}
WIDTH = {"code": 9, "level": 6, "name": 48, "desc": 40, "phase": 8, "release": 12, "prio": 8, "role": 18,
         "mode": 22, "deps": 16, "follows": 12, "effort": 11, "rate": 9, "days": 10, "start": 12, "end": 12,
         "ref": 12, "coop": 16, "note": 30}
INPUT = Font(color="0000FF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
NOTE = Font(italic=True, color="595959")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
PHASE_FILL = PatternFill("solid", fgColor="F2F2F2")
EFFORT_FMT, RATE_FMT, DAYS_FMT, FOCUS_FMT = "0.0#", "0.00", "0", "0%"
_STRFTIME = {"%d": "dd", "%m": "mm", "%Y": "yyyy", "%y": "yy", "%b": "mmm", "%B": "mmmm", "%a": "ddd", "%A": "dddd"}


@dataclass
class Layout:
    """Where verify() and the tests find things. Letters come from WBS_COLUMNS; nothing else hard-codes them."""

    wbs_sheet: str
    cols: dict[str, str]  # WBS column key -> letter
    rows: dict[str, int]  # WBS code (phase, group, package) -> row on the WBS sheet
    estimate_sheet: str
    cells: dict[str, str]  # plan_start, planned_end, working_days, total_effort -> cell on the estimate sheet
    gate_rows: dict[str, int]  # gate id -> row on the estimate sheet
    gate_date_col: str
    resources_sheet: str
    focus_col: str
    focus_rows: dict[str, int]  # role id -> row on the resources sheet


def excel_date_format(fmt: str) -> str:
    """strftime pattern of project.date_format -> Excel number format ("%d/%m/%Y" -> "dd/mm/yyyy")."""
    return re.sub(r"%[a-zA-Z]", lambda m: _STRFTIME.get(m.group(0), ""), fmt)


class _Formulas:
    """WORKDAY/NETWORKDAYS, or their .INTL forms when the weekend is not Saturday + Sunday."""

    def __init__(self, weekend: tuple[int, ...]):
        self.mask = None if set(weekend) == {5, 6} else "".join("1" if d in weekend else "0" for d in range(7))

    def workday(self, start: str, days: str) -> str:
        if self.mask is None:
            return f"WORKDAY({start},{days},HolidayDates)"
        return f'_xlfn.WORKDAY.INTL({start},{days},"{self.mask}",HolidayDates)'

    def networkdays(self, a: str, b: str) -> str:
        if self.mask is None:
            return f"NETWORKDAYS({a},{b},HolidayDates)"
        return f'_xlfn.NETWORKDAYS.INTL({a},{b},"{self.mask}",HolidayDates)'


def _ref(sheet: str, cell: str) -> str:
    return "'" + sheet.replace("'", "''") + "'!" + cell


def _abs(col: str, first: int, last: int) -> str:
    return f"${col}${first}:${col}${last}"


def _ranges(col: str, rows: list[int]) -> str:
    """Rows [7, 8, 9, 11] -> "O7:O9,O11" (consecutive rows collapse into one range)."""
    out: list[str] = []
    rows = sorted(rows)
    i = 0
    while i < len(rows):
        j = i
        while j + 1 < len(rows) and rows[j + 1] == rows[j] + 1:
            j += 1
        out.append(f"{col}{rows[i]}" if i == j else f"{col}{rows[i]}:{col}{rows[j]}")
        i = j + 1
    return ",".join(out)


def _sheet(wb: Workbook, title: str, heading: str, note: str, landscape: bool = True) -> Worksheet:
    ws = wb.create_sheet(title)
    ws["A1"], ws["A2"] = heading, note
    ws["A1"].font, ws["A2"].font = TITLE, NOTE
    if landscape:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _header(ws: Worksheet, row: int, col: int, titles: list[str]) -> None:
    for i, t in enumerate(titles):
        c = ws.cell(row, col + i, t)
        c.font, c.fill = BOLD, HEAD_FILL


def _put(ws: Worksheet, ref: str, value, fmt: str | None = None, font: Font | None = None) -> None:
    c = ws[ref]
    c.value = value
    if fmt:
        c.number_format = fmt
    if font:
        c.font = font


def _validate(ws: Worksheet, kind: str, refs: str, low: str, high: str | None = None) -> None:
    """Data validation on comma-separated cells or ranges: kind >= low, or between low and high."""
    if not refs:
        return
    dv = (DataValidation(type=kind, operator="between", formula1=low, formula2=high) if high
          else DataValidation(type=kind, operator="greaterThanOrEqual", formula1=low))
    for ref in refs.split(","):
        dv.add(ref)
    ws.add_data_validation(dv)


def build_workbook(project: Project, plan: Plan, path: Path) -> Layout:
    lang = project.languages["workbook"]

    def t(key: str, **fmt) -> str:
        return label(lang, key, **fmt)

    date_fmt = excel_date_format(project.date_format)
    f = _Formulas(project.weekend)
    names = {key: t(f"sheet_{key}") for key in ("estimate", "wbs", "resources", "assumptions", "holidays")}
    wb = Workbook()
    wb.remove(wb.active)
    est = _sheet(wb, names["estimate"], t("title_estimate", name=project.name), t("note_estimate"))
    wbs = _sheet(wb, names["wbs"], t("title_wbs"), t("note_wbs", resources=names["resources"]))
    res = _sheet(wb, names["resources"], t("title_resources"), t("note_resources"), landscape=False)
    asm = _sheet(wb, names["assumptions"], t("title_assumptions"), t("note_assumptions"))
    hol = _sheet(wb, names["holidays"], t("title_holidays"), t("note_holidays"), landscape=False)
    wb.calculation.fullCalcOnLoad = True

    # ---------------------------------------------------------------- Holidays (HolidayDates)
    _header(hol, HEADER, 1, [t("date"), t("name"), t("kind")])
    days = sorted(project.holidays.items())
    for i, (d, info) in enumerate(days):
        r = FIRST + i
        _put(hol, f"A{r}", d, date_fmt, INPUT)
        hol[f"B{r}"], hol[f"C{r}"] = info.get("name", ""), info.get("kind", "")
    last_holiday = FIRST + len(days) + SPARE_HOLIDAYS - 1
    for r in range(FIRST + len(days), last_holiday + 1):
        _put(hol, f"A{r}", None, date_fmt, INPUT)
    holidays_ref = _ref(hol.title, _abs("A", FIRST, last_holiday))
    wb.defined_names["HolidayDates"] = DefinedName("HolidayDates", attr_text=holidays_ref)
    _validate(hol, "date", f"A{FIRST}:A{last_holiday}", "1")
    hol.column_dimensions["A"].width, hol.column_dimensions["B"].width, hol.column_dimensions["C"].width = 14, 40, 16
    hol.freeze_panes = f"A{FIRST}"

    # ---------------------------------------------------------------- Resources (focus inputs)
    _header(res, HEADER, 1, [t("role_id"), t("role"), t("focus"), t("effort"), t("duties")])
    focus_rows: dict[str, int] = {}
    for i, role in enumerate(project.roles.values()):
        r = FIRST + i
        focus_rows[role.id] = r
        res[f"A{r}"], res[f"B{r}"] = role.id, project.role_title(role.id, lang)
        _put(res, f"C{r}", role.focus, FOCUS_FMT, INPUT)
        res[f"E{r}"] = text({"duties": role.duties, "i18n": role.i18n}, "duties", lang)
    res_total = FIRST + len(project.roles)
    _validate(res, "decimal", f"C{FIRST}:C{res_total - 1}", "0.01", "1")
    for w, col in zip((12, 26, 10, 14, 70), "ABCDE", strict=True):
        res.column_dimensions[col].width = w

    # ---------------------------------------------------------------- WBS
    by = plan.by_code()
    nodes = {n["code"]: n for n in [*plan.phases, *plan.groups]}
    members = {c: [p.code for p in plan.packages if p.code.startswith(c + ".")] for c in nodes}
    codes = sorted([*by, *(c for c in nodes if members[c])], key=code_key)
    rows = {code: FIRST + i for i, code in enumerate(codes)}
    last = max(FIRST, FIRST + len(codes) - 1)  # keeps column ranges valid for an empty plan
    prev = {code: (q[i - 1] if i else None) for q in role_queues(plan).values() for i, code in enumerate(q)}
    releases = {r.id: r.name for r in project.releases}
    C = COL
    _header(wbs, HEADER, 1, [t(k) for k in WBS_COLUMNS])
    inputs: dict[str, list[int]] = {"effort": [], "rate": [], "fixed": []}
    for code in codes:
        r = rows[code]
        level = len(code.split("."))
        p = by.get(code)
        item = p.as_item() if p else nodes[code]
        desc_key = "objective" if level == 1 else "desc"
        values = {"code": code, "level": level, "name": text(item, "name", lang), "desc": text(item, desc_key, lang),
                  "phase": code.split(".")[0]}
        for key, v in values.items():
            wbs[f"{C[key]}{r}"] = v
        cell = {key: f"{C[key]}{r}" for key in WBS_COLUMNS}
        if p is None:  # phase or group roll-up over its packages
            leaves = [rows[m] for m in members[code]]
            wbs[cell["start"]] = f"=MIN({_ranges(C['start'], leaves)})"
            wbs[cell["end"]] = f"=MAX({_ranges(C['end'], leaves)})"
            wbs[cell["effort"]] = f"=SUM({_ranges(C['effort'], leaves)})"
            wbs[cell["days"]] = "=" + f.networkdays(cell["start"], cell["end"])
            for key in WBS_COLUMNS:
                wbs[cell[key]].font = BOLD
                if level == 1:
                    wbs[cell[key]].fill = PHASE_FILL
            wbs.row_dimensions[r].outlineLevel = level - 1
        else:
            wbs.row_dimensions[r].outlineLevel = 2
            wbs[cell["release"]] = releases.get(p.release, p.release) if p.release else t("no_release")
            wbs[cell["prio"]] = p.prio
            wbs[cell["role"]] = project.role_title(p.role, lang)
            wbs[cell["deps"]] = ", ".join(p.deps)
            wbs[cell["follows"]] = prev.get(code) or ""
            wbs[cell["ref"]], wbs[cell["coop"]], wbs[cell["note"]] = p.ref, p.coop, p.note
            mode = p.mode
            if mode == "seq":
                focus = _ref(res.title, f"$C${focus_rows[p.role]}")
                after = ["PlanStart-1"]
                if prev.get(code):
                    after.append(f"{C['end']}{rows[prev[code]]}")
                after += [f"{C['end']}{rows[d]}" for d in p.deps]
                wbs[cell["mode"]] = t("mode_seq")
                wbs[cell["days"]] = f"=MAX(1,ROUND({cell['effort']}/{focus},0))"
                wbs[cell["start"]] = "=" + f.workday(f"MAX({','.join(after)})", "1")
                wbs[cell["end"]] = "=" + f.workday(cell["start"], f"{cell['days']}-1")
            else:
                wbs[cell["days"]] = "=" + f.networkdays(cell["start"], cell["end"])
                if mode == "fixed":
                    wbs[cell["mode"]] = t("mode_fixed")
                    _put(wbs, cell["start"], p.schedule[1], font=INPUT)
                    _put(wbs, cell["end"], p.schedule[2], font=INPUT)
                    inputs["fixed"].append(r)
                elif mode == "span":
                    a, b = p.schedule[1], p.schedule[2]
                    wbs[cell["mode"]] = t("mode_span", a=a, b=b)
                    wbs[cell["start"]] = f"={C['start']}{rows[a]}"
                    wbs[cell["end"]] = f"={C['end']}{rows[b]}"
                elif mode == "period":
                    after, months = p.schedule[1], p.schedule[2]
                    wbs[cell["mode"]] = t("mode_period", months=months, after=after)
                    wbs[cell["start"]] = "=" + f.workday(f"{C['end']}{rows[after]}", "1")
                    wbs[cell["end"]] = f"=EDATE({cell['start']},{months})-1"
                else:  # tail
                    of, n = p.schedule[1], p.schedule[2]
                    wbs[cell["mode"]] = t("mode_tail", days=n, of=of)
                    wbs[cell["end"]] = f"={C['end']}{rows[of]}"
                    wbs[cell["start"]] = "=" + f.workday(cell["end"], str(-(n - 1)))
            if p.effort_rate is not None:
                _put(wbs, cell["rate"], p.effort_rate, RATE_FMT, INPUT)
                wbs[cell["effort"]] = f"=ROUND({cell['days']}/5*{cell['rate']},1)"
                inputs["rate"].append(r)
            else:
                _put(wbs, cell["effort"], p.effort, font=INPUT)
                inputs["effort"].append(r)
        for key, fmt in (("start", date_fmt), ("end", date_fmt), ("effort", EFFORT_FMT), ("days", DAYS_FMT)):
            wbs[cell[key]].number_format = fmt
    _validate(wbs, "decimal", _ranges(C["effort"], inputs["effort"]), "0")
    _validate(wbs, "decimal", _ranges(C["rate"], inputs["rate"]), "0.001")
    fixed = [_ranges(C[k], inputs["fixed"]) for k in ("start", "end") if inputs["fixed"]]
    _validate(wbs, "date", ",".join(fixed), "1")
    wbs.sheet_properties.outlinePr.summaryBelow = False
    wbs.freeze_panes = f"{C['name']}{FIRST}"
    for key in WBS_COLUMNS:
        wbs.column_dimensions[C[key]].width = WIDTH[key]

    def W(key: str) -> str:  # absolute WBS column range for SUMIF(S)
        return _ref(wbs.title, _abs(C[key], FIRST, last))

    for r in focus_rows.values():
        _put(res, f"D{r}", f"=SUMIF({W('role')},B{r},{W('effort')})", EFFORT_FMT)
    res[f"A{res_total}"] = t("total")
    _put(res, f"D{res_total}", f"=SUM(D{FIRST}:D{res_total - 1})", EFFORT_FMT, BOLD)
    res[f"A{res_total}"].font = BOLD

    # ---------------------------------------------------------------- Assumptions
    _header(asm, HEADER, 1, [t("id"), t("kind"), t("assumption"), t("scope"), t("source")])
    for i, a in enumerate(plan.assumptions):
        r = FIRST + i
        for col, key in zip("ABCDE", ("id", "kind", "text", "scope", "source"), strict=True):
            asm[f"{col}{r}"] = text(a, key, lang) if key != "id" else a["id"]
    for w, col in zip((8, 16, 70, 16, 24), "ABCDE", strict=True):
        asm.column_dimensions[col].width = w

    # ---------------------------------------------------------------- Estimate
    cells = {"plan_start": "B4", "planned_end": "B5", "working_days": "B6", "total_effort": "B7"}
    est["A4"], est["A5"], est["A6"], est["A7"] = t("plan_start"), t("planned_end"), t("working_days"), t("total_effort")
    _put(est, "B4", plan.start, date_fmt, INPUT)
    _validate(est, "date", "B4", "1")
    wb.defined_names["PlanStart"] = DefinedName("PlanStart", attr_text=_ref(est.title, "$B$4"))
    _put(est, "B5", f"=MAX({W('end')})", date_fmt)
    _put(est, "B6", "=" + f.networkdays(f"MIN({W('start')})", "B5"), DAYS_FMT)
    _put(est, "B7", f"=SUMIF({W('level')},3,{W('effort')})", EFFORT_FMT)
    for r in range(4, 8):
        est[f"A{r}"].font = BOLD

    # milestones
    r = 9
    est[f"A{r}"], est[f"A{r}"].font = t("milestones"), BOLD
    r += 1
    _header(est, r, 1, [t("gate"), t("name"), t("wbs"), t("planned_date"), t("hard_date"), t("spare_days"),
                        t("verdict"), t("criteria")])
    gate_rows: dict[str, int] = {}
    for g in plan.gates:
        r += 1
        gate_rows[g["id"]] = r
        est[f"A{r}"], est[f"B{r}"], est[f"C{r}"] = g["id"], text(g, "name", lang), g["at"]
        est[f"H{r}"] = text(g, "criteria", lang)
        if g["at"] in rows:
            _put(est, f"D{r}", "=" + _ref(wbs.title, f"${C['end']}${rows[g['at']]}"), date_fmt)
        _put(est, f"E{r}", project.hard_dates.get(g["id"]), date_fmt, INPUT)
        spare = (f'=IF(OR(D{r}="",E{r}=""),"",IF(D{r}<=E{r},{f.networkdays(f"D{r}+1", f"E{r}")},'
                 f'-{f.networkdays(f"E{r}+1", f"D{r}")}))')
        _put(est, f"F{r}", spare, DAYS_FMT)
        est[f"G{r}"] = f'=IF(F{r}="","",IF(F{r}<0,"{t("late")}","{t("on_time")}"))'
    if gate_rows:
        first_gate, last_gate = min(gate_rows.values()), max(gate_rows.values())
        _validate(est, "date", f"E{first_gate}:E{last_gate}", "1")
        est.conditional_formatting.add(f"F{first_gate}:G{last_gate}",
                                       FormulaRule(formula=[f"$F{first_gate}<0"], font=Font(color="C00000")))

    # effort by phase and role
    roles = list(project.roles)
    rc = {rid: get_column_letter(3 + i) for i, rid in enumerate(roles)}
    total_col, start_col, end_col, days_col = (get_column_letter(3 + len(roles) + i) for i in range(4))
    r += 2
    est[f"A{r}"], est[f"A{r}"].font = t("by_phase"), BOLD
    r += 1
    head = r
    _header(est, head, 1, [t("phase"), t("name"), *(project.role_title(x, lang) for x in roles), t("total"),
                           t("start"), t("end"), t("working_days")])
    phase_rows = []
    for ph in plan.phases:
        if ph["code"] not in rows:
            continue
        r += 1
        phase_rows.append(r)
        w = rows[ph["code"]]
        est[f"A{r}"], est[f"B{r}"] = ph["code"], text(ph, "name", lang)
        for rid in roles:
            col = rc[rid]
            _put(est, f"{col}{r}", f"=SUMIFS({W('effort')},{W('role')},{col}${head},{W('phase')},$A{r})",
                 EFFORT_FMT)
        _put(est, f"{total_col}{r}", f"=SUM({rc[roles[0]]}{r}:{rc[roles[-1]]}{r})", EFFORT_FMT)
        _put(est, f"{start_col}{r}", "=" + _ref(wbs.title, f"${C['start']}${w}"), date_fmt)
        _put(est, f"{end_col}{r}", "=" + _ref(wbs.title, f"${C['end']}${w}"), date_fmt)
        _put(est, f"{days_col}{r}", "=" + _ref(wbs.title, f"${C['days']}${w}"), DAYS_FMT)
    r += 1
    est[f"A{r}"], est[f"A{r}"].font = t("total"), BOLD
    if phase_rows:
        a, b = phase_rows[0], phase_rows[-1]
        for col in [*rc.values(), total_col]:
            _put(est, f"{col}{r}", f"=SUM({col}{a}:{col}{b})", EFFORT_FMT, BOLD)
        _put(est, f"{start_col}{r}", f"=MIN({start_col}{a}:{start_col}{b})", date_fmt, BOLD)
        _put(est, f"{end_col}{r}", f"=MAX({end_col}{a}:{end_col}{b})", date_fmt, BOLD)
        _put(est, f"{days_col}{r}", "=" + f.networkdays(f"{start_col}{r}", f"{end_col}{r}"), DAYS_FMT, BOLD)

    # effort by release and role
    r += 2
    est[f"A{r}"], est[f"A{r}"].font = t("by_release"), BOLD
    r += 1
    head = r
    _header(est, head, 1, [t("release"), t("name"), *(project.role_title(x, lang) for x in roles), t("total")])
    release_rows = [(rel.id, rel.name) for rel in project.releases]
    if any(not p.release for p in plan.packages):
        release_rows.append(("", t("no_release")))
    first_rel = r + 1
    for rid_, name in release_rows:
        r += 1
        est[f"A{r}"], est[f"B{r}"] = rid_, name
        for rid in roles:
            col = rc[rid]
            _put(est, f"{col}{r}", f"=SUMIFS({W('effort')},{W('role')},{col}${head},{W('release')},$B{r})",
                 EFFORT_FMT)
        _put(est, f"{total_col}{r}", f"=SUM({rc[roles[0]]}{r}:{rc[roles[-1]]}{r})", EFFORT_FMT)
    r += 1
    est[f"A{r}"], est[f"A{r}"].font = t("total"), BOLD
    for col in [*rc.values(), total_col]:
        _put(est, f"{col}{r}", f"=SUM({col}{first_rel}:{col}{r - 1})", EFFORT_FMT, BOLD)
    est.column_dimensions["A"].width, est.column_dimensions["B"].width = 26, 44
    for i in range(3, 9 + len(roles)):
        est.column_dimensions[get_column_letter(i)].width = 14

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return Layout(wbs_sheet=wbs.title, cols=dict(COL), rows=rows, estimate_sheet=est.title, cells=cells,
                  gate_rows=gate_rows, gate_date_col="D", resources_sheet=res.title, focus_col="C",
                  focus_rows=focus_rows)
