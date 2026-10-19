"""Recalculate the estimate workbook with LibreOffice in Docker and compare it with the baseline schedule."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .errors import KickoffError
from .model import Plan
from .schedule import Schedule

# Shipped inside the package so it is found after `uv tool install <plugin>/cli` (code runs from site-packages).
DOCKERFILE_DIR = Path(__file__).resolve().parent / "docker" / "lo-calc"
IMAGE = "kickoff-lo-calc"
ERRORS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")


def docker_available() -> bool:
    if not shutil.which("docker"):
        return False
    return subprocess.run(["docker", "info"], capture_output=True).returncode == 0


def ensure_image(image: str = IMAGE) -> None:
    if subprocess.run(["docker", "image", "inspect", image], capture_output=True).returncode == 0:
        return
    r = subprocess.run(["docker", "build", "-t", image, str(DOCKERFILE_DIR)], capture_output=True, text=True,
                       encoding="utf-8")
    if r.returncode:
        raise KickoffError(f"docker build failed: {r.stderr[-500:]}")


def recalc(path: Path, image: str = IMAGE) -> Path:
    """A recalculated copy of `path` (LibreOffice headless --convert-to xlsx); the input is not changed."""
    ensure_image(image)
    work = Path(tempfile.mkdtemp(prefix="kickoff-lo-"))
    shutil.copy(path, work / "in.xlsx")
    r = subprocess.run(["docker", "run", "--rm", "-v", f"{work}:/w", image, "soffice", "--headless", "--calc",
                        "--convert-to", "xlsx", "--outdir", "/w/out", "/w/in.xlsx"],
                       capture_output=True, text=True, encoding="utf-8")
    out = work / "out" / "in.xlsx"
    if r.returncode or not out.exists():
        raise KickoffError(f"LibreOffice recalculation failed: {r.stderr[-500:]}")
    return out


def _as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def verify(path: Path, layout, sched: Schedule, plan: Plan) -> list[str]:
    """Problems in a recalculated workbook: error cells, and every WBS row or gate that differs from `sched`."""
    wb = load_workbook(path, data_only=True)
    problems = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRORS):
                    problems.append(f"{ws.title}!{c.coordinate}: {c.value}")
    ws, col = wb[layout.wbs_sheet], layout.cols
    for code, r in layout.rows.items():
        slot = sched.slots.get(code)
        if not slot:
            problems.append(f"{code}: in the workbook but not in the schedule")
            continue
        s, e = _as_date(ws[f"{col['start']}{r}"].value), _as_date(ws[f"{col['end']}{r}"].value)
        days, effort = ws[f"{col['days']}{r}"].value, ws[f"{col['effort']}{r}"].value
        if (s, e, days) != (slot.start, slot.end, slot.workdays):
            problems.append(f"{code}: workbook {s} – {e} ({days} d), schedule {slot.start} – {slot.end} "
                            f"({slot.workdays} d)")
        if not isinstance(effort, int | float) or abs(effort - slot.effort) > 0.005:
            problems.append(f"{code}: workbook effort {effort}, schedule {slot.effort}")
    gates = sched.gate_dates(plan)
    es = wb[layout.estimate_sheet]
    for gid, r in layout.gate_rows.items():
        v = _as_date(es[f"{layout.gate_date_col}{r}"].value)
        if gid in gates and v != gates[gid]:
            problems.append(f"{gid}: workbook {v}, schedule {gates[gid]}")
    total = sum(sched.slots[p.code].effort for p in plan.packages)
    got = es[layout.cells["total_effort"]].value
    if not isinstance(got, int | float) or abs(got - total) > 0.005:
        problems.append(f"total effort: workbook {got}, schedule {total:g}")
    return problems
