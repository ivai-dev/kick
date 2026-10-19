import json
from dataclasses import replace
from pathlib import Path

import pytest
from helpers import pkg, plan, project, write_kickoff
from openpyxl import load_workbook

from kickoff import lo_verify
from kickoff.cli import main
from kickoff.labels import LABELS
from kickoff.model import load_plan, load_project
from kickoff.schedule import build_schedule
from kickoff.workbook import WBS_COLUMNS, build_workbook

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
SRC = Path(__file__).resolve().parents[1] / "src" / "kickoff"
MODES = [
    pkg("1.1.1", effort=2),
    pkg("1.1.2", role="qa", deps=["1.1.1"]),
    pkg("1.1.3", role="lead", schedule={"fixed": ["2026-09-24", "2026-09-25"]}),
    pkg("1.1.4", role="lead", effort=0, schedule={"span": ["1.1.1", "1.1.2"]}, effort_rate=0.2),
    pkg("1.1.5", role="lead", schedule={"period": {"after": "1.1.2", "months": 12}}),
    pkg("1.1.6", role="lead", schedule={"tail": {"of": "1.1.5", "days": 5}}),
    pkg("1.1.7", effort=1),
]


def formulas(path):
    wb = load_workbook(path)
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")}


def values(path):
    wb = load_workbook(path)
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row}


def build_example(tmp_path, lang=None):
    root = EXAMPLES / "ai-minute"
    pr, pl = load_project(root), load_plan(root)
    if lang:
        pr.languages = {**pr.languages, "workbook": lang}
    out = tmp_path / "ai-minute.xlsx"
    return pr, pl, out, build_workbook(pr, pl, out)


def test_formula_shapes_for_every_mode(tmp_path):
    out = tmp_path / "modes.xlsx"
    lay = build_workbook(project(), plan(MODES), out)
    ws = load_workbook(out)[lay.wbs_sheet]
    r, c = lay.rows, lay.cols

    def cell(key, code):
        return ws[f"{c[key]}{r[code]}"].value

    focus_dev = f"'{lay.resources_sheet}'!$C${lay.focus_rows['dev']}"
    assert cell("start", "1.1.1") == "=WORKDAY(MAX(PlanStart-1),1,HolidayDates)"
    assert cell("days", "1.1.1") == f"=MAX(1,ROUND({c['effort']}{r['1.1.1']}/{focus_dev},0))"
    assert cell("end", "1.1.1") == f"=WORKDAY({c['start']}{r['1.1.1']},{c['days']}{r['1.1.1']}-1,HolidayDates)"
    assert cell("start", "1.1.2") == f"=WORKDAY(MAX(PlanStart-1,{c['end']}{r['1.1.1']}),1,HolidayDates)"
    assert cell("start", "1.1.7") == f"=WORKDAY(MAX(PlanStart-1,{c['end']}{r['1.1.1']}),1,HolidayDates)"
    assert cell("follows", "1.1.7") == "1.1.1"
    assert not isinstance(cell("start", "1.1.3"), str)  # fixed dates are inputs
    assert cell("start", "1.1.4") == f"={c['start']}{r['1.1.1']}"
    assert cell("end", "1.1.4") == f"={c['end']}{r['1.1.2']}"
    assert cell("effort", "1.1.4") == f"=ROUND({c['days']}{r['1.1.4']}/5*{c['rate']}{r['1.1.4']},1)"
    assert cell("start", "1.1.5") == f"=WORKDAY({c['end']}{r['1.1.2']},1,HolidayDates)"
    assert cell("end", "1.1.5") == f"=EDATE({c['start']}{r['1.1.5']},12)-1"
    assert cell("end", "1.1.6") == f"={c['end']}{r['1.1.5']}"
    assert cell("start", "1.1.6") == f"=WORKDAY({c['end']}{r['1.1.6']},-4,HolidayDates)"
    leaves = f"{c['start']}{r['1.1.1']}:{c['start']}{r['1.1.7']}"
    assert cell("start", "1.1") == f"=MIN({leaves})"
    assert cell("start", "1") == f"=MIN({leaves})"
    assert cell("days", "1.1.3") == f"=NETWORKDAYS({c['start']}{r['1.1.3']},{c['end']}{r['1.1.3']},HolidayDates)"


def test_other_weekend_uses_intl_functions(tmp_path):
    out = tmp_path / "fri-sat.xlsx"
    pr = project(calendar={"weekend": ["fri", "sat"], "holidays": []})
    lay = build_workbook(pr, plan([pkg("1.1.1")]), out)
    start = load_workbook(out)[lay.wbs_sheet][f"{lay.cols['start']}{lay.rows['1.1.1']}"].value
    assert start == '=_xlfn.WORKDAY.INTL(MAX(PlanStart-1),1,"0000110",HolidayDates)'


def test_example_sheets_names_and_size(tmp_path):
    pr, pl, out, lay = build_example(tmp_path)
    wb = load_workbook(out)
    vi = LABELS["vi"]
    assert wb.sheetnames == [vi[f"sheet_{k}"] for k in ("estimate", "wbs", "resources", "assumptions", "holidays")]
    assert {"HolidayDates", "PlanStart"} <= set(wb.defined_names)
    assert set(lay.rows) == set(build_schedule(pr, pl, baseline=True).slots)
    assert set(lay.gate_rows) == {g["id"] for g in pl.gates}
    header = [c.value for c in wb[lay.wbs_sheet][4]][: len(WBS_COLUMNS)]
    assert header == [vi[k] for k in WBS_COLUMNS]
    assert 400 < len(formulas(out)) < 1000  # about 4 per WBS row plus the summaries; the old design had ~11,000


def test_english_labels_and_sheet_name_rules(tmp_path):
    _, _, out, _ = build_example(tmp_path, lang="en")
    assert load_workbook(out).sheetnames == ["Estimate", "WBS", "Resources", "Assumptions", "Holidays"]
    for lang in LABELS:
        for key in LABELS[lang]:
            if key.startswith("sheet_"):
                name = LABELS[lang][key]
                assert len(name) <= 31 and not set(name) & set("[]:*?/\\"), name
        assert set(LABELS[lang]) == set(LABELS["en"]), lang


def test_tracking_fields_do_not_change_the_workbook(tmp_path):
    tracked = [dict(p, done=0.5, status="done") if p["code"] == "1.1.1" else p for p in MODES]
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    build_workbook(project(), plan(MODES), a)
    build_workbook(project(), plan(tracked), b)
    assert values(a) == values(b)


def test_example_baseline_equals_golden_schedule():
    root = EXAMPLES / "ai-minute"
    expected = json.loads((root / "expected" / "schedule.json").read_text(encoding="utf-8"))
    sched = build_schedule(load_project(root), load_plan(root), baseline=True)
    for code, (start, end, workdays, *_rest) in expected.items():
        s = sched.slots[code]
        assert (s.start.isoformat(), s.end.isoformat(), s.workdays) == (start, end, workdays), code


def test_workbook_has_no_project_literals():
    for name in ("workbook.py", "labels.py"):
        src = (SRC / name).read_text(encoding="utf-8")
        for word in ("AI Minute", "ASR", "GĐ", "TDA", "LTV", "KTV"):
            assert word not in src, f"{name}: {word}"


def test_command_writes_default_path(tmp_path, capsys):
    write_kickoff(tmp_path)
    assert main(["--root", str(tmp_path), "workbook"]) == 0
    assert (tmp_path / "docs" / "06-planning" / "demo-estimate.xlsx").exists()
    assert "wrote" in capsys.readouterr().out


def test_dockerfile_ships_inside_the_package():
    assert (lo_verify.DOCKERFILE_DIR / "Dockerfile").is_file()
    assert lo_verify.DOCKERFILE_DIR.is_relative_to(SRC)


@pytest.mark.docker
@pytest.mark.skipif(not lo_verify.docker_available(), reason="Docker not available")
def test_libreoffice_recalculation_matches_schedule(tmp_path):
    pr, pl, out, lay = build_example(tmp_path)
    assert lo_verify.verify(lo_verify.recalc(out), lay, build_schedule(pr, pl, baseline=True), pl) == []


@pytest.mark.docker
@pytest.mark.skipif(not lo_verify.docker_available(), reason="Docker not available")
def test_editing_focus_and_effort_moves_the_dates(tmp_path):
    pr, pl = project(), plan(MODES)
    out = tmp_path / "edit.xlsx"
    lay = build_workbook(pr, pl, out)
    wb = load_workbook(out)
    wb[lay.resources_sheet][f"{lay.focus_col}{lay.focus_rows['dev']}"] = 0.5
    wb[lay.wbs_sheet][f"{lay.cols['effort']}{lay.rows['1.1.7']}"] = 3
    wb.save(out)
    edited = plan([dict(p, effort=3) if p["code"] == "1.1.7" else p for p in MODES])
    pr.roles["dev"] = replace(pr.roles["dev"], focus=0.5)
    assert lo_verify.verify(lo_verify.recalc(out), lay, build_schedule(pr, edited, baseline=True), edited) == []
